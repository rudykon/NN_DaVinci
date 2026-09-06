#!/usr/bin/env python3
"""Generate and independently inspect the 0.6.1 publication artifact corpus.

This stage deliberately does not consume :func:`figure_proof` as a pass gate.
Its persisted SVG outputs are inputs to ``figure_svg_oracle.mjs`` in the full
controller, after every requested format has been written.
"""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
from typing import Any
import xml.etree.ElementTree as ET
import zipfile

from nn_davinci.figure_export import (
    PNG_EXPORT_DPI,
    compile_figure,
    export_figure,
    provenance_report,
)
from nn_davinci.figure_templates import TEMPLATE_SPECS, instantiate_template
from nn_davinci.model_figure import model_figure_from_graph, validate_model_figure_provenance
from nn_davinci.project import Project
from nn_davinci.real_models import import_real_model, real_model_registry
from nn_davinci.semantic import derive_semantic_view


VERSION = "0.6.1"
RELEASE = "0.6.1 Beta — Publication Fidelity & Model-to-Figure Completion"
FORMATS = ("svg", "pdf", "tikz", "png", "eps", "pptx", "html")
SIGNATURES = {
    ".svg": b"<?xml",
    ".pdf": b"%PDF",
    ".tex": b"\\documentclass",
    ".png": b"\x89PNG\r\n\x1a\n",
    ".eps": b"%!PS-Adobe-3.0 EPSF-3.0",
    ".pptx": b"PK",
    ".html": b"<!doctype html>",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run(arguments: list[str], *, cwd: Path | None = None, timeout: int = 180) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        arguments,
        cwd=cwd,
        check=False,
        capture_output=True,
        text=True,
        timeout=timeout,
    )


def pdf_fonts(path: Path) -> tuple[int, bool, str]:
    result = run(["pdffonts", str(path)])
    if result.returncode:
        return 0, False, (result.stderr or result.stdout).strip()
    lines = [line for line in result.stdout.splitlines()[2:] if line.strip() and not set(line.strip()) <= {"-"}]
    if not lines:
        return 0, False, "pdffonts reported no font rows"
    header = result.stdout.splitlines()[0] if result.stdout.splitlines() else ""
    try:
        embedded_column = header.index("emb")
    except ValueError:
        return len(lines), False, "pdffonts output lacks an emb column"
    embedded = [line[embedded_column : embedded_column + 3].strip() == "yes" for line in lines]
    return len(lines), all(embedded), ""


def compile_tikz(source: Path, destination: Path) -> tuple[bool, str]:
    with tempfile.TemporaryDirectory(prefix="nndv-061-tikz-") as directory:
        temporary = Path(directory)
        staged = temporary / "figure.tex"
        shutil.copy2(source, staged)
        result = run(
            [
                "pdflatex",
                "-no-shell-escape",
                "-interaction=nonstopmode",
                "-halt-on-error",
                "-output-directory",
                str(temporary),
                str(staged),
            ],
            timeout=180,
        )
        compiled = temporary / "figure.pdf"
        if result.returncode or not compiled.is_file():
            return False, (result.stdout + "\n" + result.stderr)[-1_200:]
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(compiled, destination)
    fonts, embedded, detail = pdf_fonts(destination)
    if fonts == 0 or not embedded:
        return False, detail or "compiled TikZ PDF has non-embedded fonts"
    return True, ""


def png_metadata(path: Path) -> tuple[dict[str, Any], list[str]]:
    failures: list[str] = []
    try:
        from PIL import Image

        with Image.open(path) as image:
            width, height = image.size
            dpi = tuple(float(value) for value in image.info.get("dpi", (0.0, 0.0)))
            mode = image.mode
            embedded = json.loads(image.info.get("nndv.figure_export", "{}"))
    except Exception as exc:
        return {}, [f"PNG inspection failed: {type(exc).__name__}: {exc}"]
    if width not in {2102, 2103} or height != 1394:
        failures.append(f"PNG size is {width}x{height}, expected 2102/2103x1394 for 178x118 mm at 300 DPI")
    if mode != "RGBA":
        failures.append(f"PNG mode is {mode}, expected RGBA")
    if not dpi or any(abs(value - PNG_EXPORT_DPI) > 0.1 for value in dpi[:2]):
        failures.append(f"PNG DPI is {dpi}, expected {PNG_EXPORT_DPI}")
    if embedded.get("dpi") != PNG_EXPORT_DPI:
        failures.append("PNG lacks the explicit NN_DaVinci 300-DPI metadata record")
    return {"width_px": width, "height_px": height, "dpi": dpi, "mode": mode}, failures


def pptx_editability(path: Path, page_count: int) -> tuple[dict[str, Any], list[str]]:
    failures: list[str] = []
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        presentation = Presentation(path)
        shapes = [shape for slide in presentation.slides for shape in slide.shapes]
        pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes)
        texts = sum(bool(getattr(shape, "has_text_frame", False)) for shape in shapes)
        connectors = sum(shape.name.startswith("NNDV editable connector") for shape in shapes)
        tensor_faces = sum(shape.name.startswith("NNDV tensor face") for shape in shapes)
        panels = sum(shape.name.startswith("NNDV Panel ") for shape in shapes)
        slides = len(presentation.slides)
        if slides != page_count:
            failures.append(f"PPTX has {slides} slides for {page_count} Figure pages")
        if pictures:
            failures.append(f"PPTX contains {pictures} raster picture shape(s)")
        if not shapes or not texts or not panels:
            failures.append("PPTX lacks editable shapes, text, or Panel objects")
        with zipfile.ZipFile(path) as archive:
            slide_xml = b"".join(
                archive.read(name)
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            )
        if b"NNDV editable" not in slide_xml:
            failures.append("PPTX OOXML lacks NN_DaVinci editable object markers")
        return {
            "slides": slides,
            "shapes": len(shapes),
            "texts": texts,
            "connectors": connectors,
            "tensor_faces": tensor_faces,
            "panels": panels,
            "pictures": pictures,
        }, failures
    except Exception as exc:
        return {}, [f"PPTX inspection failed: {type(exc).__name__}: {exc}"]


def inspect_outputs(paths: list[Path], *, page_count: int) -> tuple[dict[str, Any], list[str], list[Path]]:
    failures: list[str] = []
    by_suffix: dict[str, list[Path]] = {}
    for path in paths:
        by_suffix.setdefault(path.suffix.lower(), []).append(path)
        signature = SIGNATURES.get(path.suffix.lower())
        if not path.is_file() or path.stat().st_size == 0:
            failures.append(f"missing or empty export: {path.name}")
        elif signature is None or not path.read_bytes().startswith(signature):
            failures.append(f"invalid {path.suffix} signature: {path.name}")
    observed_formats = {
        "svg" if suffix == ".svg" else
        "tikz" if suffix == ".tex" else
        suffix.lstrip(".")
        for suffix in by_suffix
    }
    if observed_formats != set(FORMATS):
        failures.append(f"format set is {sorted(observed_formats)}, expected {sorted(FORMATS)}")

    svg_details: list[dict[str, Any]] = []
    for path in by_suffix.get(".svg", []):
        try:
            root = ET.parse(path).getroot()
            metadata = next((item for item in root.iter() if item.tag.rsplit("}", 1)[-1] == "metadata" and item.get("id") == "nndv-figure-ir"), None)
            document = json.loads(metadata.text) if metadata is not None and metadata.text else {}
            if document.get("nn_davinci_version") != VERSION:
                failures.append(f"SVG metadata version is not {VERSION}: {path.name}")
            svg_details.append({"file": path.name, "metadata_version": document.get("nn_davinci_version")})
        except (OSError, ET.ParseError, json.JSONDecodeError) as exc:
            failures.append(f"SVG parse failed for {path.name}: {exc}")

    pdf_details: list[dict[str, Any]] = []
    for path in by_suffix.get(".pdf", []):
        count, embedded, detail = pdf_fonts(path)
        if not embedded:
            failures.append(f"PDF fonts are not fully embedded for {path.name}: {detail}")
        pdf_details.append({"file": path.name, "font_count": count, "all_embedded": embedded})

    compiled_pdfs: list[Path] = []
    tikz_details: list[dict[str, Any]] = []
    for path in by_suffix.get(".tex", []):
        destination = path.with_suffix(".tikz.pdf")
        passed, tikz_detail = compile_tikz(path, destination)
        if not passed:
            failures.append(f"TikZ compile failed for {path.name}: {tikz_detail}")
        else:
            compiled_pdfs.append(destination)
        tikz_details.append({"file": path.name, "compiled_pdf": destination.name, "passed": passed})

    png_details: list[dict[str, Any]] = []
    proof_paths: list[Path] = []
    for path in by_suffix.get(".png", []):
        png_detail, png_failures = png_metadata(path)
        failures.extend(f"{path.name}: {message}" for message in png_failures)
        proof = path.with_name(f"{path.stem}.proof-300dpi.png")
        shutil.copy2(path, proof)
        proof_paths.append(proof)
        png_details.append({"file": path.name, "proof": proof.name, **png_detail})

    eps_details: list[dict[str, Any]] = []
    for path in by_suffix.get(".eps", []):
        payload = path.read_bytes()
        vector = b"%%NNDVVectorPolicy: SVG-to-PDF-to-Poppler-EPS" in payload
        fonts = b"%%BeginResource: font" in payload
        if not vector or not fonts:
            failures.append(f"EPS lacks vector/font evidence: {path.name}")
        eps_details.append({"file": path.name, "vector": vector, "font_resources": fonts})

    html_details: list[dict[str, Any]] = []
    for path in by_suffix.get(".html", []):
        source = path.read_text(encoding="utf-8")
        external_asset_tokens = ('src="http://', 'src="https://', 'href="http://', 'href="https://')
        offline = "default-src 'none'" in source and not any(token in source for token in external_asset_tokens)
        editable = "Download edited page SVG" in source and 'id="nndv-figure-ir"' in source
        if not offline or not editable:
            failures.append(f"HTML is not a self-contained editable offline artifact: {path.name}")
        html_details.append({"file": path.name, "offline": offline, "editable": editable})

    pptx_details: list[dict[str, Any]] = []
    for path in by_suffix.get(".pptx", []):
        pptx_detail, pptx_failures = pptx_editability(path, page_count)
        failures.extend(f"{path.name}: {message}" for message in pptx_failures)
        pptx_details.append({"file": path.name, **pptx_detail})
    return {
        "svg": svg_details,
        "pdf": pdf_details,
        "tikz": tikz_details,
        "png": png_details,
        "eps": eps_details,
        "pptx": pptx_details,
        "html": html_details,
    }, failures, [*compiled_pdfs, *proof_paths]


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def generate_one(kind: str, key: str, destination: Path) -> dict[str, Any]:
    destination.mkdir(parents=True, exist_ok=False)
    failures: list[str] = []
    if kind == "template":
        graph, figure = instantiate_template(key)
        semantic = derive_semantic_view(graph)
        provenance_validation = {"passed": True, "scope": "bundled-template"}
    else:
        # The faithful module view is the release-reviewed projection path.  It
        # keeps the complete imported node/edge/port/tensor inventory in Graph
        # IR while model_figure_from_graph produces the bounded paper view.
        graph = import_real_model(key, view="module")
        semantic = derive_semantic_view(graph)
        figure = model_figure_from_graph(
            graph,
            semantic_view=semantic,
            level="stage",
            view="paper",
            mode="mixed",
            page_preset="double-column",
        )
        provenance_validation = validate_model_figure_provenance(figure, graph, semantic)
        if not provenance_validation.get("passed"):
            failures.append("Graph IR ↔ Semantic View ↔ Figure IR provenance validation failed")
    figure.validate()
    compiled = compile_figure(figure)
    paths = export_figure(figure, destination / "figure", formats=FORMATS)
    inspection, inspection_failures, derived_paths = inspect_outputs(paths, page_count=len(figure.pages))
    failures.extend(inspection_failures)

    project = Project(
        name=figure.name,
        graph=graph,
        model_source={
            "kind": kind,
            "key": key,
            "offline": True,
            "weights_downloaded": False,
        },
        semantic_view={
            "version": "1.0",
            "level": "template" if kind == "template" else "stage",
            "view": "paper",
            "document": semantic.to_dict(),
        },
        export={"formats": list(FORMATS), "transparent": False},
        figure_ir=figure.to_dict(),
    )
    project_path = project.save(destination / "figure.nndv.json")
    roundtrip = Project.load(project_path)
    persisted_semantic = roundtrip.persisted_semantic_view()
    semantic_roundtrip_exact = not (
        roundtrip.graph.to_dict() != graph.to_dict()
        or roundtrip.figure_ir != figure.to_dict()
        or persisted_semantic is None
        or persisted_semantic.to_dict() != semantic.to_dict()
    )
    if not semantic_roundtrip_exact:
        failures.append(".nndv.json Graph/Figure round-trip differs from generated evidence")
    persisted_semantic_ids = {
        entity.id for entity in persisted_semantic.entities
    } if persisted_semantic is not None else set()
    figure_semantic_source_ids = {
        item.provenance.source_id
        for item in figure.iter_objects()
        if item.provenance.kind == "semantic_view"
    }
    provenance_path = destination / "provenance.json"
    provenance = provenance_report(figure)
    provenance["semantic_view"] = semantic.to_dict()
    provenance["model_figure_validation"] = provenance_validation
    write_json(provenance_path, provenance)

    all_paths = [*paths, *derived_paths, project_path, provenance_path]
    files = {
        path.relative_to(destination).as_posix(): {
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        }
        for path in sorted(all_paths)
    }
    export_suffixes = Counter(path.suffix.lower() for path in paths)
    return {
        "kind": kind,
        "key": key,
        "status": "PASS" if not failures else "FAIL",
        "figure_digest": figure.digest(),
        "graph_digest": hashlib.sha256(
            json.dumps(graph.to_dict(), ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
        ).hexdigest(),
        "page_count": len(figure.pages),
        "panel_count": sum(len(page.panels) for page in figure.pages),
        "object_count": sum(1 for _ in figure.iter_objects()),
        "compiled_primitive_count": len(compiled),
        "requested_formats": list(FORMATS),
        "export_file_suffix_counts": dict(sorted(export_suffixes.items())),
        "inspection": inspection,
        "provenance_validation": provenance_validation,
        "project_semantic_view": {
            "field": "semantic_view.document",
            "source_digest": semantic.source_digest,
            "entity_count": len(semantic.entities),
            "figure_semantic_source_id_count": len(figure_semantic_source_ids),
            "orphan_figure_semantic_source_ids": sorted(
                figure_semantic_source_ids - persisted_semantic_ids
            ),
            "roundtrip_exact": semantic_roundtrip_exact,
        },
        "release_gate_uses_python_figure_proof": False,
        "files": files,
        "failures": failures,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    output_dir = args.output_dir.resolve()
    if output_dir.exists() and any(output_dir.iterdir()):
        raise SystemExit(f"refusing to overwrite non-empty publication output directory: {output_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)

    results: list[dict[str, Any]] = []
    generation_failures: list[dict[str, str]] = []
    jobs = [
        *(('template', spec.slug) for spec in TEMPLATE_SPECS),
        *(('real-model', key) for key in real_model_registry()),
    ]
    for kind, key in jobs:
        destination = output_dir / ("templates" if kind == "template" else "real-models") / key
        try:
            result = generate_one(kind, key, destination)
        except Exception as exc:
            result = {
                "kind": kind,
                "key": key,
                "status": "FAIL",
                "requested_formats": list(FORMATS),
                "release_gate_uses_python_figure_proof": False,
                "files": {},
                "failures": [f"{type(exc).__name__}: {exc}"],
            }
        results.append(result)
        generation_failures.extend(
            {"kind": kind, "key": key, "message": message}
            for message in result.get("failures", [])
        )

    generated_files = [
        path
        for path in output_dir.rglob("*")
        if path.is_file() and not path.is_symlink()
    ]
    svg_files = sorted(path for path in generated_files if path.suffix.lower() == ".svg")
    counts = {
        "figures": len(results),
        "templates": sum(item["kind"] == "template" for item in results),
        "real_models": sum(item["kind"] == "real-model" for item in results),
        "requested_formats": len(results) * len(FORMATS),
        "svg_files": len(svg_files),
        "tikz_compiled_pdfs": sum(path.name.endswith(".tikz.pdf") for path in generated_files),
        "proof_pngs_300dpi": sum(path.name.endswith(".proof-300dpi.png") for path in generated_files),
        "project_files": sum(path.name == "figure.nndv.json" for path in generated_files),
        "provenance_files": sum(path.name == "provenance.json" for path in generated_files),
        "failures": len(generation_failures),
    }
    expected_counts = {
        "figures": 14,
        "templates": 7,
        "real_models": 7,
        "requested_formats": 98,
        "svg_files": 14,
        "tikz_compiled_pdfs": 14,
        "proof_pngs_300dpi": 14,
        "project_files": 14,
        "provenance_files": 14,
        "failures": 0,
    }
    for name, expected in expected_counts.items():
        if counts[name] != expected:
            generation_failures.append({"kind": "corpus", "key": name, "message": f"count {counts[name]} != {expected}"})
    counts["failures"] = len(generation_failures)
    report = {
        "schema_version": "nndv-0.6.1-publication-artifact-corpus-1",
        "release": RELEASE,
        "status": "PASS" if not generation_failures else "FAIL",
        "fresh_outputs": True,
        "formats": list(FORMATS),
        "strict_svg_oracle_required_after_generation": True,
        "python_figure_proof_is_release_blocker": False,
        "output_root": str(output_dir),
        "counts": counts,
        "expected_counts": expected_counts,
        "svg_inputs": [
            {
                "relative_path": path.relative_to(output_dir).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
            for path in svg_files
        ],
        "results": results,
        "failures": generation_failures,
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print("NNDV_061_PUBLICATION_ARTIFACTS=" + json.dumps({"status": report["status"], **counts}, sort_keys=True))
    return 0 if not generation_failures else 1


if __name__ == "__main__":
    raise SystemExit(main())
