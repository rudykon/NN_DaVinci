"""Genuine vector, print-proof and offline exports for Scene IR 1.0.

All two-dimensional formats consume :class:`ProjectedScene` directly.  Scene
IR is never converted to Figure IR and no browser/WebGL screenshot participates
in publication export.
"""

from __future__ import annotations

import html
import json
import math
from pathlib import Path
import re
import zipfile
import zlib
from typing import Any, Iterable, Mapping
import xml.etree.ElementTree as ET

from .errors import ExportError, OptionalDependencyError
from .pdf_fonts import EmbeddedTrueTypeFont, load_scene_pdf_fonts, select_font_runs
from .scene_gltf import (
    export_scene_glb,
    export_scene_gltf,
    export_scene_json,
    validate_scene_assets,
)
from .scene_projection import PT_TO_MM, ProjectedPrimitive, ProjectedScene, ProjectionOptions, project_scene
from .version import __version__


SCENE_EXPORT_VERSION = "nndv-scene-export-1"
PNG_EXPORT_DPI = 300
MM_TO_PT = 72.0 / 25.4
MM_TO_EMU = 36_000


def _projection(
    scene_or_projection: Any,
    *,
    camera_id: str | None,
    options: ProjectionOptions | Mapping[str, Any] | None,
) -> ProjectedScene:
    if isinstance(scene_or_projection, ProjectedScene):
        if camera_id is not None or options is not None:
            raise ExportError("A pre-projected scene cannot be combined with camera_id or projection options")
        scene_or_projection.validate()
        return scene_or_projection
    validate_scene_assets(scene_or_projection)
    publication_options = options if options is not None else ProjectionOptions(
        auto_frame=True,
        density="paper",
        target_occupancy=0.72,
    )
    return project_scene(scene_or_projection, camera_id=camera_id, options=publication_options)


def _esc(value: Any) -> str:
    return html.escape(str(value), quote=True)


def _format(value: float) -> str:
    result = f"{float(value):.6f}".rstrip("0").rstrip(".")
    return result if result not in {"", "-0"} else "0"


def _style(primitive: ProjectedPrimitive) -> tuple[str, str, float, float, str]:
    style = primitive.style
    fill = str(style.get("fill", "none"))
    stroke = str(style.get("stroke", "#334155"))
    stroke_width = float(style.get("stroke_width_pt", 0.75))
    opacity = max(0.0, min(1.0, float(style.get("opacity", 1.0))))
    dash = str(style.get("dash", ""))
    return fill, stroke, stroke_width, opacity, dash


def _semantic_svg_attributes(primitive: ProjectedPrimitive) -> str:
    attributes: list[str] = []
    mapping = {
        "architecture_role_id": "data-architecture-role-id",
        "architecture_role": "data-architecture-role",
        "architecture_route_id": "data-architecture-route-id",
        "architecture_route_role": "data-architecture-route-role",
        "architecture_role_graph_edge_id": "data-role-graph-edge-id",
        "architecture_direction": "data-architecture-direction",
        "architecture_evidence_digest": "data-evidence-digest",
        "architecture_role_graph_digest": "data-role-graph-digest",
        "repeat_count": "data-repeat-count",
    }
    for key, attribute in mapping.items():
        value = primitive.metadata.get(key)
        if value is not None and value != "":
            attributes.append(f'{attribute}="{_esc(value)}"')
    endpoint_roles = primitive.metadata.get("endpoint_role_ids", [])
    if isinstance(endpoint_roles, list) and endpoint_roles:
        attributes.append(
            'data-endpoint-role-ids="'
            + _esc(json.dumps(endpoint_roles, ensure_ascii=False, separators=(",", ":")))
            + '"'
        )
    return " ".join(attributes)


def _presentation_binding(primitive: ProjectedPrimitive, projected: ProjectedScene) -> dict[str, Any]:
    return {
        "object_id": primitive.object_id,
        "depth": primitive.depth,
        "kind": primitive.kind,
        "source_digest": projected.source_digest,
        "architecture_role_id": primitive.metadata.get("architecture_role_id"),
        "architecture_role": primitive.metadata.get("architecture_role"),
        "architecture_route_id": primitive.metadata.get("architecture_route_id"),
        "architecture_route_role": primitive.metadata.get("architecture_route_role"),
        "architecture_role_graph_edge_id": primitive.metadata.get("architecture_role_graph_edge_id"),
        "architecture_direction": primitive.metadata.get("architecture_direction"),
        "architecture_evidence_digest": primitive.metadata.get("architecture_evidence_digest"),
        "architecture_role_graph_digest": primitive.metadata.get("architecture_role_graph_digest"),
        "endpoint_object_ids": primitive.metadata.get("endpoint_object_ids", []),
        "endpoint_role_ids": primitive.metadata.get("endpoint_role_ids", []),
        "repeat_count": primitive.metadata.get("repeat_count", 1),
    }


def _arrow_head(points: list[tuple[float, float]], stroke_width_pt: float) -> list[tuple[float, float]]:
    if len(points) < 2:
        return []
    end, before = points[-1], points[-2]
    dx, dy = end[0] - before[0], end[1] - before[1]
    length = math.hypot(dx, dy)
    if length <= 1.0e-9:
        return []
    ux, uy = dx / length, dy / length
    size = max(1.25, stroke_width_pt * PT_TO_MM * 4.5)
    width = size * 0.55
    base = (end[0] - ux * size, end[1] - uy * size)
    return [end, (base[0] - uy * width, base[1] + ux * width), (base[0] + uy * width, base[1] - ux * width)]


def _bezier_samples(points: list[tuple[float, float]], count: int = 49) -> list[tuple[float, float]]:
    if len(points) != 4:
        return points
    result: list[tuple[float, float]] = []
    for index in range(count):
        amount = index / (count - 1)
        inverse = 1.0 - amount
        weights = (inverse**3, 3.0 * inverse**2 * amount, 3.0 * inverse * amount**2, amount**3)
        result.append(tuple(sum(weights[item] * points[item][axis] for item in range(4)) for axis in range(2)))  # type: ignore[arg-type]
    return result


def render_scene_svg(
    scene_or_projection: Any,
    *,
    camera_id: str | None = None,
    options: ProjectionOptions | Mapping[str, Any] | None = None,
) -> str:
    """Return a standalone SVG whose geometry remains native vector DOM."""

    projected = _projection(scene_or_projection, camera_id=camera_id, options=options)
    metadata = json.dumps(projected.to_dict(), ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    lines = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        (
            f'<svg xmlns="http://www.w3.org/2000/svg" width="{_format(projected.width_mm)}mm" '
            f'height="{_format(projected.height_mm)}mm" viewBox="0 0 {_format(projected.width_mm)} {_format(projected.height_mm)}" '
            f'data-nndv-export="{SCENE_EXPORT_VERSION}" data-nndv-version="{_esc(__version__)}" '
            f'data-camera-id="{_esc(projected.camera_id)}">'
        ),
        f'<metadata id="nndv-scene-projection">{_esc(metadata)}</metadata>',
        "<desc>Deterministic CPU vector projection of NN_DaVinci Scene IR; no WebGL screenshot.</desc>",
        f'<rect id="nndv-background" x="0" y="0" width="{_format(projected.width_mm)}" height="{_format(projected.height_mm)}" fill="{_esc(projected.options.background)}"/>',
    ]
    for index, primitive in enumerate(projected.primitives):
        fill, stroke, stroke_width, opacity, dash = _style(primitive)
        attributes = (
            f'id="nndv-{primitive.kind}-{index}" data-kind="{primitive.kind}" '
            f'data-object-id="{_esc(primitive.object_id)}" data-depth="{_format(primitive.depth)}" '
            f'stroke="{_esc(stroke)}" stroke-width="{_format(stroke_width * PT_TO_MM)}" '
            f'vector-effect="non-scaling-stroke" opacity="{_format(opacity)}"'
        )
        semantic_attributes = _semantic_svg_attributes(primitive)
        if semantic_attributes:
            attributes += " " + semantic_attributes
        endpoint_ids = primitive.metadata.get("endpoint_object_ids", [])
        if isinstance(endpoint_ids, list) and endpoint_ids:
            attributes += f' data-endpoint-object-ids="{_esc(json.dumps(endpoint_ids, ensure_ascii=False, separators=(",", ":")))}"'
        if dash:
            attributes += f' stroke-dasharray="{_esc(dash)}"'
        point_text = " ".join(f"{_format(x)},{_format(y)}" for x, y in primitive.points)
        if primitive.kind == "face":
            lines.append(f'<polygon {attributes} points="{point_text}" fill="{_esc(fill)}" stroke-linejoin="round"/>')
        elif primitive.kind in {"edge", "polyline", "leader"}:
            lines.append(f'<polyline {attributes} points="{point_text}" fill="none" stroke-linecap="round" stroke-linejoin="round"/>')
        elif primitive.kind == "arrow":
            lines.append(f'<polyline {attributes} points="{point_text}" fill="none" stroke-linecap="round" stroke-linejoin="round"/>')
            head = _arrow_head(primitive.points, stroke_width)
            head_text = " ".join(f"{_format(x)},{_format(y)}" for x, y in head)
            if head:
                lines.append(
                    f'<polygon data-kind="arrowhead" data-object-id="{_esc(primitive.object_id)}" points="{head_text}" fill="{_esc(stroke)}" stroke="none" opacity="{_format(opacity)}"/>'
                )
        elif primitive.kind == "bezier" and len(primitive.points) == 4:
            start, first, second, end = primitive.points
            path = f"M {_format(start[0])} {_format(start[1])} C {_format(first[0])} {_format(first[1])}, {_format(second[0])} {_format(second[1])}, {_format(end[0])} {_format(end[1])}"
            lines.append(f'<path {attributes} d="{path}" fill="none" stroke-linecap="round"/>')
        elif primitive.kind == "label" and primitive.points:
            x, y = primitive.points[0]
            font_size = max(7.0, float(primitive.style.get("font_size_pt", 7.0)))
            font_size_mm = math.ceil(font_size * PT_TO_MM * 1_000_000.0) / 1_000_000.0
            bbox = primitive.style.get("text_bbox_mm", [x, y, x + 10.0, y + font_size * PT_TO_MM * 1.25])
            lines.append(
                f'<rect data-kind="label-halo" data-object-id="{_esc(primitive.object_id)}" x="{_format(float(bbox[0]) - 0.35)}" '
                f'y="{_format(float(bbox[1]) - 0.2)}" width="{_format(float(bbox[2]) - float(bbox[0]) + 0.7)}" '
                f'height="{_format(float(bbox[3]) - float(bbox[1]) + 0.4)}" rx="0.5" fill="{_esc(projected.options.background)}" fill-opacity="0.82" stroke="none"/>'
            )
            baseline = y + font_size * PT_TO_MM
            lines.append(
                f'<text id="nndv-label-{index}" data-kind="label" data-object-id="{_esc(primitive.object_id)}" '
                f'data-depth="{_format(primitive.depth)}" x="{_format(x)}" y="{_format(baseline)}" '
                f'{semantic_attributes} '
                f'fill="{_esc(fill)}" opacity="{_format(opacity)}" stroke="none" '
                f'font-family="{_esc(primitive.style.get("font_family", projected.options.font_family))}" '
                f'data-font-size-pt="{_format(font_size)}" font-size="{_format(font_size_mm)}" '
                f'xml:space="preserve">{_esc(primitive.text)}</text>'
            )
    lines.append("</svg>")
    payload = "\n".join(lines) + "\n"
    validate_scene_svg(payload)
    return payload


def validate_scene_svg(payload: str) -> None:
    try:
        root = ET.fromstring(payload)
    except ET.ParseError as exc:
        raise ExportError(f"Scene SVG export is not well-formed XML: {exc}") from exc
    if root.tag.rsplit("}", 1)[-1] != "svg" or root.get("data-nndv-export") != SCENE_EXPORT_VERSION:
        raise ExportError("Scene SVG export is missing its vector export marker")
    prohibited = {"image", "foreignObject", "script"}
    for element in root.iter():
        if element.tag.rsplit("}", 1)[-1] in prohibited:
            raise ExportError("Scene SVG contains a raster, executable, or foreign-object element")
        for key, value in element.attrib.items():
            if key.rsplit("}", 1)[-1] in {"href", "src"} and not value.startswith("#"):
                raise ExportError("Scene SVG contains an external resource reference")
    text = payload.lower()
    if "javascript:" in text or "data:text/html" in text:
        raise ExportError("Scene SVG contains an executable URI")


def _tex_escape(value: Any) -> str:
    replacements = {
        "\\": r"\textbackslash{}",
        "{": r"\{",
        "}": r"\}",
        "$": r"\$",
        "&": r"\&",
        "#": r"\#",
        "%": r"\%",
        "_": r"\_",
        "^": r"\textasciicircum{}",
        "~": r"\textasciitilde{}",
        # Keep common scientific punctuation in the selected text font.  Raw
        # UTF-8 under pdfLaTeX can otherwise become an unmapped Type-3 glyph,
        # breaking copy/paste and independent text extraction.
        "×": r"\texttimes{}",
        "·": r"\textperiodcentered{}",
    }
    return "".join(replacements.get(character, character) for character in str(value))


def _tikz_colour(value: str) -> str:
    text = str(value).lstrip("#")
    if len(text) != 6 or any(character not in "0123456789abcdefABCDEF" for character in text):
        text = "334155"
    red, green, blue = (int(text[index : index + 2], 16) for index in (0, 2, 4))
    return f"{{rgb,255:red,{red};green,{green};blue,{blue}}}"


def render_scene_tikz(
    scene_or_projection: Any,
    *,
    camera_id: str | None = None,
    options: ProjectionOptions | Mapping[str, Any] | None = None,
) -> str:
    projected = _projection(scene_or_projection, camera_id=camera_id, options=options)
    lines = [
        r"\documentclass{article}",
        r"\usepackage[T1]{fontenc}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage{lmodern}",
        r"\usepackage{textcomp}",
        r"\usepackage{tikz}",
        r"\usetikzlibrary{arrows.meta}",
        rf"\usepackage[paperwidth={projected.width_mm:g}mm,paperheight={projected.height_mm:g}mm,margin=0mm,noheadfoot]{{geometry}}",
        r"\pagestyle{empty}",
        r"\setlength{\parindent}{0pt}",
        r"\begin{document}",
        r"\begin{tikzpicture}[remember picture,overlay,x=1mm,y=-1mm,shift={(current page.north west)}]",
        rf"\path[fill={_tikz_colour(projected.options.background)}] (0,0) rectangle ({projected.width_mm:g},{projected.height_mm:g});",
        rf"% NN_DaVinci {__version__}; {SCENE_EXPORT_VERSION}; camera={_tex_escape(projected.camera_id)}; digest={projected.source_digest}",
    ]
    for primitive in projected.primitives:
        binding = _presentation_binding(primitive, projected)
        if binding.get("architecture_role_id") or binding.get("architecture_route_id"):
            lines.append("% NNDV-PRESENTATION-BINDING " + json.dumps(binding, ensure_ascii=True, sort_keys=True, separators=(",", ":")))
        fill, stroke, stroke_width, opacity, dash = _style(primitive)
        draw_options = f"draw={_tikz_colour(stroke)},line width={stroke_width:g}pt,opacity={opacity:g}"
        if dash:
            draw_options += ",dashed"
        points = " -- ".join(f"({_format(x)},{_format(y)})" for x, y in primitive.points)
        if primitive.kind == "face":
            lines.append(rf"\path[{draw_options},fill={_tikz_colour(fill)},line join=round] {points} -- cycle;")
        elif primitive.kind in {"edge", "polyline", "leader"}:
            lines.append(rf"\draw[{draw_options},line cap=round,line join=round] {points};")
        elif primitive.kind == "arrow":
            lines.append(rf"\draw[-{{Latex[length=2.2mm,width=1.2mm]}},{draw_options},line cap=round] {points};")
        elif primitive.kind == "bezier" and len(primitive.points) == 4:
            start, first, second, end = primitive.points
            lines.append(
                rf"\draw[{draw_options},line cap=round] ({_format(start[0])},{_format(start[1])}) .. controls "
                rf"({_format(first[0])},{_format(first[1])}) and ({_format(second[0])},{_format(second[1])}) .. "
                rf"({_format(end[0])},{_format(end[1])});"
            )
        elif primitive.kind == "label" and primitive.points:
            x, y = primitive.points[0]
            font_size = max(7.0, float(primitive.style.get("font_size_pt", 7.0)))
            lines.append(
                rf"\node[anchor=north west,inner sep=0.4mm,fill={_tikz_colour(projected.options.background)},fill opacity=.82,text opacity=1,text={_tikz_colour(fill)},"
                rf"font=\sffamily\fontsize{{{font_size:g}}}{{{font_size * 1.2:g}}}\selectfont] at ({_format(x)},{_format(y)}) "
                rf"{{{_tex_escape(primitive.text)}}};"
            )
    lines.extend([r"\end{tikzpicture}\null", r"\end{document}"])
    payload = "\n".join(lines) + "\n"
    if r"\includegraphics" in payload or r"\write18" in payload or r"\begin{tikzpicture}" not in payload:
        raise ExportError("TikZ Scene export failed its native-vector validation")
    return payload


def _rgb(value: str) -> tuple[float, float, float]:
    text = str(value).lstrip("#")
    if len(text) != 6 or any(character not in "0123456789abcdefABCDEF" for character in text):
        text = "334155"
    return tuple(int(text[index : index + 2], 16) / 255.0 for index in (0, 2, 4))  # type: ignore[return-value]


def _pdf_text(value: str) -> str:
    encoded = value.encode("latin-1", "replace").decode("latin-1")
    return encoded.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _pdf_content(
    projected: ProjectedScene,
    fonts: tuple[EmbeddedTrueTypeFont, ...],
) -> tuple[bytes, dict[str, dict[int, str]]]:
    height_pt = projected.height_mm * MM_TO_PT
    commands: list[str] = ["q"]
    used: dict[str, dict[int, str]] = {font.resource_name: {} for font in fonts}
    background = _rgb(projected.options.background)
    commands.append(f"{background[0]:.6f} {background[1]:.6f} {background[2]:.6f} rg 0 0 {projected.width_mm * MM_TO_PT:.6f} {height_pt:.6f} re f")

    def coordinate(point: tuple[float, float]) -> tuple[float, float]:
        return point[0] * MM_TO_PT, height_pt - point[1] * MM_TO_PT

    for primitive in projected.primitives:
        binding = _presentation_binding(primitive, projected)
        if binding.get("architecture_role_id") or binding.get("architecture_route_id"):
            commands.append("% NNDV-PRESENTATION-BINDING " + json.dumps(binding, ensure_ascii=True, sort_keys=True, separators=(",", ":")))
        fill, stroke, stroke_width, _opacity, _dash = _style(primitive)
        stroke_rgb, fill_rgb = _rgb(stroke), _rgb(fill)
        commands.extend(
            [
                f"{stroke_rgb[0]:.6f} {stroke_rgb[1]:.6f} {stroke_rgb[2]:.6f} RG",
                f"{fill_rgb[0]:.6f} {fill_rgb[1]:.6f} {fill_rgb[2]:.6f} rg",
                f"{stroke_width:.6f} w",
            ]
        )
        if primitive.kind in {"face", "edge", "polyline", "leader", "arrow"} and primitive.points:
            start = coordinate(primitive.points[0])
            commands.append(f"{start[0]:.6f} {start[1]:.6f} m")
            for point in primitive.points[1:]:
                x, y = coordinate(point)
                commands.append(f"{x:.6f} {y:.6f} l")
            if primitive.kind == "face":
                commands.append("h B")
            else:
                commands.append("S")
            if primitive.kind == "arrow":
                head = _arrow_head(primitive.points, stroke_width)
                if head:
                    x, y = coordinate(head[0])
                    commands.append(f"{x:.6f} {y:.6f} m")
                    for point in head[1:]:
                        x, y = coordinate(point)
                        commands.append(f"{x:.6f} {y:.6f} l")
                    commands.append("h f")
        elif primitive.kind == "bezier" and len(primitive.points) == 4:
            start, first, second, end = [coordinate(point) for point in primitive.points]
            commands.append(f"{start[0]:.6f} {start[1]:.6f} m {first[0]:.6f} {first[1]:.6f} {second[0]:.6f} {second[1]:.6f} {end[0]:.6f} {end[1]:.6f} c S")
        elif primitive.kind == "label" and primitive.points:
            x, y = coordinate(primitive.points[0])
            font_size = max(7.0, float(primitive.style.get("font_size_pt", 7.0)))
            commands.append(f"BT {x:.6f} {y - font_size:.6f} Td")
            for font, glyphs, source_text in select_font_runs(primitive.text, fonts):
                for glyph_id, character in zip(glyphs, source_text, strict=True):
                    used[font.resource_name].setdefault(glyph_id, character)
                encoded = "".join(f"{glyph_id:04X}" for glyph_id in glyphs)
                commands.append(f"/{font.resource_name} {font_size:.6f} Tf <{encoded}> Tj")
            commands.append("ET")
    commands.append("Q")
    return ("\n".join(commands) + "\n").encode("ascii"), used


def _pdf_stream(data: bytes, *, compress: bool = False, extra: str = "") -> bytes:
    original_length = len(data)
    payload = zlib.compress(data, level=9) if compress else data
    filters = " /Filter /FlateDecode" if compress else ""
    length1 = f" /Length1 {original_length}" if compress else ""
    suffix = f" {extra.strip()}" if extra.strip() else ""
    return (
        f"<< /Length {len(payload)}{length1}{filters}{suffix} >>\nstream\n".encode("ascii")
        + payload
        + b"\nendstream"
    )


def _to_unicode_cmap(font: EmbeddedTrueTypeFont, used: Mapping[int, str]) -> bytes:
    mappings = sorted(used.items())
    lines = [
        "/CIDInit /ProcSet findresource begin",
        "12 dict begin",
        "begincmap",
        "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def",
        f"/CMapName /NNDV-{font.resource_name}-UCS def",
        "/CMapType 2 def",
        "1 begincodespacerange",
        "<0000> <FFFF>",
        "endcodespacerange",
    ]
    for start in range(0, len(mappings), 100):
        chunk = mappings[start : start + 100]
        lines.append(f"{len(chunk)} beginbfchar")
        for glyph_id, character in chunk:
            unicode_hex = character.encode("utf-16-be").hex().upper()
            lines.append(f"<{glyph_id:04X}> <{unicode_hex}>")
        lines.append("endbfchar")
    lines.extend([
        "endcmap",
        "CMapName currentdict /CMap defineresource pop",
        "end",
        "end",
    ])
    return ("\n".join(lines) + "\n").encode("ascii")


def _pdf_font_objects(
    objects: list[bytes | None],
    font: EmbeddedTrueTypeFont,
    used: Mapping[int, str],
) -> int:
    type0_reference = len(objects) + 1
    cid_reference = type0_reference + 1
    descriptor_reference = type0_reference + 2
    file_reference = type0_reference + 3
    unicode_reference = type0_reference + 4
    base = font.postscript_name
    bbox = " ".join(str(font.metric_1000(value)) for value in font.bbox)
    widths = " ".join(
        f"{glyph_id} [{font.width_1000(glyph_id)}]"
        for glyph_id in sorted(used)
    )
    objects.extend([
        (
            f"<< /Type /Font /Subtype /Type0 /BaseFont /{base} /Encoding /Identity-H "
            f"/DescendantFonts [{cid_reference} 0 R] /ToUnicode {unicode_reference} 0 R >>"
        ).encode("ascii"),
        (
            f"<< /Type /Font /Subtype /CIDFontType2 /BaseFont /{base} "
            "/CIDSystemInfo << /Registry (Adobe) /Ordering (Identity) /Supplement 0 >> "
            f"/FontDescriptor {descriptor_reference} 0 R /DW 1000 /W [{widths}] /CIDToGIDMap /Identity >>"
        ).encode("ascii"),
        (
            f"<< /Type /FontDescriptor /FontName /{base} /Flags 32 /FontBBox [{bbox}] "
            f"/ItalicAngle 0 /Ascent {font.metric_1000(font.ascent)} "
            f"/Descent {font.metric_1000(font.descent)} /CapHeight {font.metric_1000(font.ascent)} "
            f"/StemV 80 /MissingWidth 1000 /FontFile2 {file_reference} 0 R >>"
        ).encode("ascii"),
        _pdf_stream(font.data, compress=True),
        _pdf_stream(_to_unicode_cmap(font, used), compress=True),
    ])
    return type0_reference


def render_scene_pdf(
    scene_or_projection: Any,
    *,
    camera_id: str | None = None,
    options: ProjectionOptions | Mapping[str, Any] | None = None,
) -> bytes:
    """Return a self-contained vector PDF using native path/text operators."""

    projected = _projection(scene_or_projection, camera_id=camera_id, options=options)
    fonts = load_scene_pdf_fonts()
    content, used_glyphs = _pdf_content(projected, fonts)
    width_pt, height_pt = projected.width_mm * MM_TO_PT, projected.height_mm * MM_TO_PT
    objects: list[bytes | None] = [None, None, None, _pdf_stream(content)]
    font_references: dict[str, int] = {}
    for font in fonts:
        used = used_glyphs[font.resource_name]
        if used:
            font_references[font.resource_name] = _pdf_font_objects(objects, font, used)
    if not font_references:
        raise ExportError("Scene PDF contains no label glyphs to bind to an embedded font")
    info_reference = len(objects) + 1
    objects.append((
            f"<< /Producer (NN_DaVinci deterministic Scene vector exporter) /Title (Scene {projected.scene_id}) "
            f"/Subject (NN_DaVinci {__version__}; {SCENE_EXPORT_VERSION}; embedded-packaged-fonts; camera={_pdf_text(projected.camera_id)}; digest={projected.source_digest}) >>"
        ).encode("latin-1", "replace"))
    font_resources = " ".join(
        f"/{name} {reference} 0 R" for name, reference in font_references.items()
    )
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>"
    objects[2] = (
        f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width_pt:.6f} {height_pt:.6f}] "
        f"/Resources << /Font << {font_resources} >> >> /Contents 4 0 R >>"
    ).encode("ascii")
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, value in enumerate(objects, 1):
        if value is None:
            raise ExportError(f"Internal Scene PDF object {index} was not materialized")
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(value)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R /Info {info_reference} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii"))
    payload = bytes(output)
    validate_scene_pdf(payload)
    return payload


def validate_scene_pdf(payload: bytes) -> None:
    if not payload.startswith(b"%PDF-") or not payload.rstrip().endswith(b"%%EOF") or b"xref" not in payload:
        raise ExportError("Scene PDF export has an invalid PDF structure")
    if re.search(rb"/Subtype\s*/Image\b", payload) or b"/ImageMask" in payload:
        raise ExportError("Scene PDF export contains a raster image object")
    if not any(operator in payload for operator in (b" m\n", b" re f", b" Tj")):
        raise ExportError("Scene PDF export contains no native vector/text operators")
    if b"/BaseFont /Helvetica" in payload or b"/FontFile2" not in payload or b"/ToUnicode" not in payload:
        raise ExportError("Scene PDF must use packaged embedded TrueType fonts with Unicode maps; Base-14 fallback is forbidden")


def render_scene_eps(
    scene_or_projection: Any,
    *,
    camera_id: str | None = None,
    options: ProjectionOptions | Mapping[str, Any] | None = None,
) -> bytes:
    """Return an EPSF Level 3 vector program without a raster fallback."""

    projected = _projection(scene_or_projection, camera_id=camera_id, options=options)
    width_pt, height_pt = projected.width_mm * MM_TO_PT, projected.height_mm * MM_TO_PT
    lines = [
        "%!PS-Adobe-3.0 EPSF-3.0",
        f"%%BoundingBox: 0 0 {math.ceil(width_pt)} {math.ceil(height_pt)}",
        f"%%HiResBoundingBox: 0 0 {width_pt:.6f} {height_pt:.6f}",
        f"%%Title: NN_DaVinci Scene {projected.scene_id}",
        "%%Creator: NN_DaVinci deterministic Scene vector exporter",
        f"%%NNDVVectorPolicy: {SCENE_EXPORT_VERSION}; native PostScript paths; no raster fallback",
        f"%%NNDVCamera: {projected.camera_id}",
        f"%%NNDVSourceDigest: {projected.source_digest}",
        "%%LanguageLevel: 3",
        "%%Pages: 1",
        "%%EndComments",
    ]

    def coordinate(point: tuple[float, float]) -> tuple[float, float]:
        return point[0] * MM_TO_PT, height_pt - point[1] * MM_TO_PT

    background = _rgb(projected.options.background)
    lines.append(f"{background[0]:.6f} {background[1]:.6f} {background[2]:.6f} setrgbcolor 0 0 {width_pt:.6f} {height_pt:.6f} rectfill")
    for primitive in projected.primitives:
        fill, stroke, stroke_width, _opacity, dash = _style(primitive)
        stroke_rgb, fill_rgb = _rgb(stroke), _rgb(fill)
        lines.append(f"{stroke_width:.6f} setlinewidth")
        lines.append("[2 1] 0 setdash" if dash else "[] 0 setdash")
        if primitive.kind in {"face", "edge", "polyline", "leader", "arrow"} and primitive.points:
            x, y = coordinate(primitive.points[0])
            path = [f"newpath {x:.6f} {y:.6f} moveto"]
            for point in primitive.points[1:]:
                x, y = coordinate(point)
                path.append(f"{x:.6f} {y:.6f} lineto")
            if primitive.kind == "face":
                lines.extend(
                    path
                    + [
                        "closepath gsave",
                        f"{fill_rgb[0]:.6f} {fill_rgb[1]:.6f} {fill_rgb[2]:.6f} setrgbcolor fill grestore",
                        f"{stroke_rgb[0]:.6f} {stroke_rgb[1]:.6f} {stroke_rgb[2]:.6f} setrgbcolor stroke",
                    ]
                )
            else:
                lines.extend(path + [f"{stroke_rgb[0]:.6f} {stroke_rgb[1]:.6f} {stroke_rgb[2]:.6f} setrgbcolor stroke"])
            if primitive.kind == "arrow":
                head = _arrow_head(primitive.points, stroke_width)
                if head:
                    x, y = coordinate(head[0])
                    lines.append(f"newpath {x:.6f} {y:.6f} moveto")
                    for point in head[1:]:
                        x, y = coordinate(point)
                        lines.append(f"{x:.6f} {y:.6f} lineto")
                    lines.append("closepath fill")
        elif primitive.kind == "bezier" and len(primitive.points) == 4:
            start, first, second, end = [coordinate(point) for point in primitive.points]
            lines.extend(
                [
                    f"{stroke_rgb[0]:.6f} {stroke_rgb[1]:.6f} {stroke_rgb[2]:.6f} setrgbcolor",
                    f"newpath {start[0]:.6f} {start[1]:.6f} moveto {first[0]:.6f} {first[1]:.6f} {second[0]:.6f} {second[1]:.6f} {end[0]:.6f} {end[1]:.6f} curveto stroke",
                ]
            )
        elif primitive.kind == "label" and primitive.points:
            x, y = coordinate(primitive.points[0])
            font_size = max(7.0, float(primitive.style.get("font_size_pt", 7.0)))
            lines.extend(
                [
                    f"{fill_rgb[0]:.6f} {fill_rgb[1]:.6f} {fill_rgb[2]:.6f} setrgbcolor",
                    f"/Helvetica findfont {font_size:.6f} scalefont setfont",
                    f"{x:.6f} {y - font_size:.6f} moveto ({_pdf_text(primitive.text)}) show",
                ]
            )
    lines.extend(["showpage", "%%EOF"])
    payload = ("\n".join(lines) + "\n").encode("latin-1", "replace")
    validate_scene_eps(payload)
    return payload


def validate_scene_eps(payload: bytes) -> None:
    if not payload.startswith(b"%!PS-Adobe-3.0 EPSF-3.0") or b"%%NNDVVectorPolicy:" not in payload:
        raise ExportError("Scene EPS export is not a marked EPSF Level 3 vector file")
    operators = bytearray()
    string_depth = 0
    escaped = False
    comment = False
    for value in payload:
        if comment:
            if value in (10, 13):
                comment = False
                operators.append(value)
            continue
        if string_depth:
            if escaped:
                escaped = False
            elif value == 92:  # PostScript string escape: backslash
                escaped = True
            elif value == 40:  # nested opening parenthesis
                string_depth += 1
            elif value == 41:
                string_depth -= 1
            if string_depth == 0:
                operators.append(32)
            continue
        if value == 37:  # comment marker: percent
            comment = True
        elif value == 40:
            string_depth = 1
            operators.append(32)
        else:
            operators.append(value)
    lowered = bytes(operators).lower()
    if re.search(rb"\b(?:colorimage|imagemask|image)\b", lowered):
        raise ExportError("Scene EPS export contains a raster image operator")
    if b"lineto" not in lowered and b"curveto" not in lowered and b" show" not in lowered and b"rectfill" not in lowered:
        raise ExportError("Scene EPS export contains no native vector/text operators")


def export_scene_png(
    scene_or_projection: Any,
    destination: str | Path,
    *,
    camera_id: str | None = None,
    options: ProjectionOptions | Mapping[str, Any] | None = None,
) -> Path:
    """Write a direct 300-DPI CPU raster proof of the vector projection."""

    try:
        from PIL import Image, ImageDraw, ImageFont, PngImagePlugin
    except ImportError as exc:
        raise OptionalDependencyError("Scene PNG proof export requires Pillow") from exc
    projected = _projection(scene_or_projection, camera_id=camera_id, options=options)
    width_px = max(1, round(projected.width_mm / 25.4 * PNG_EXPORT_DPI))
    height_px = max(1, round(projected.height_mm / 25.4 * PNG_EXPORT_DPI))
    scale = PNG_EXPORT_DPI / 25.4
    image = Image.new("RGBA", (width_px, height_px), projected.options.background)
    draw = ImageDraw.Draw(image, "RGBA")

    def points(values: list[tuple[float, float]]) -> list[tuple[int, int]]:
        return [(round(x * scale), round(y * scale)) for x, y in values]

    for primitive in projected.primitives:
        fill, stroke, stroke_width, opacity, _dash = _style(primitive)
        alpha = round(opacity * 255)
        fill_rgba = (*[round(value * 255) for value in _rgb(fill)], alpha)
        stroke_rgba = (*[round(value * 255) for value in _rgb(stroke)], alpha)
        width = max(1, round(stroke_width / 72.0 * PNG_EXPORT_DPI))
        rendered_points = points(primitive.points)
        if primitive.kind == "face" and len(rendered_points) >= 3:
            draw.polygon(rendered_points, fill=fill_rgba, outline=stroke_rgba, width=width)
        elif primitive.kind in {"edge", "polyline", "leader", "arrow"} and len(rendered_points) >= 2:
            draw.line(rendered_points, fill=stroke_rgba, width=width, joint="curve")
            if primitive.kind == "arrow":
                head = points(_arrow_head(primitive.points, stroke_width))
                if head:
                    draw.polygon(head, fill=stroke_rgba)
        elif primitive.kind == "bezier":
            draw.line(points(_bezier_samples(primitive.points)), fill=stroke_rgba, width=width, joint="curve")
        elif primitive.kind == "label" and rendered_points:
            font_px = max(1, round(max(7.0, float(primitive.style.get("font_size_pt", 7.0))) / 72.0 * PNG_EXPORT_DPI))
            font = None
            for candidate in (
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            ):
                try:
                    font = ImageFont.truetype(candidate, font_px)
                    break
                except OSError:
                    continue
            if font is None:
                font = ImageFont.load_default()
            draw.text(rendered_points[0], primitive.text, fill=fill_rgba, font=font)
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("Software", "NN_DaVinci deterministic Scene CPU projector")
    metadata.add_text(
        "nndv.scene_export",
        json.dumps(
            {
                "schema_version": SCENE_EXPORT_VERSION,
                "source_digest": projected.source_digest,
                "camera_id": projected.camera_id,
                "dpi": PNG_EXPORT_DPI,
                "pixel_width": width_px,
                "pixel_height": height_px,
                "source": "cpu-vector-projection",
            },
            sort_keys=True,
            separators=(",", ":"),
        ),
    )
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="PNG", dpi=(PNG_EXPORT_DPI, PNG_EXPORT_DPI), pnginfo=metadata)
    with Image.open(path) as check:
        dpi = check.info.get("dpi", (0.0, 0.0))
        if check.size != (width_px, height_px) or min(dpi) < 299.0:
            raise ExportError("Scene PNG proof did not retain its required 300-DPI dimensions/metadata")
    return path


def _pptx_colour(value: str) -> Any:
    from pptx.dml.color import RGBColor

    text = str(value).lstrip("#")
    if len(text) != 6 or any(character not in "0123456789abcdefABCDEF" for character in text):
        text = "334155"
    return RGBColor.from_string(text.upper())


def export_scene_pptx(
    scene_or_projection: Any,
    destination: str | Path,
    *,
    camera_id: str | None = None,
    options: ProjectionOptions | Mapping[str, Any] | None = None,
) -> Path:
    """Write editable Office DrawingML shapes, never a flattened picture."""

    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Emu, Mm, Pt
    except ImportError as exc:
        raise OptionalDependencyError("Editable Scene PowerPoint export requires python-pptx") from exc
    projected = _projection(scene_or_projection, camera_id=camera_id, options=options)
    presentation = Presentation()
    presentation.slide_width = Mm(projected.width_mm)
    presentation.slide_height = Mm(projected.height_mm)
    presentation.core_properties.title = f"NN_DaVinci Scene {projected.scene_id}"
    presentation.core_properties.subject = "Editable deterministic CPU projection of Scene IR 1.0"
    presentation.core_properties.comments = (
        f"NN_DaVinci {__version__};{SCENE_EXPORT_VERSION};camera={projected.camera_id};digest={projected.source_digest}"
    )
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Mm(0), Mm(0), Mm(projected.width_mm), Mm(projected.height_mm))
    background.name = "NNDV Scene background"
    background.fill.solid()
    background.fill.fore_color.rgb = _pptx_colour(projected.options.background)
    background.line.fill.background()
    for index, primitive in enumerate(projected.primitives):
        fill, stroke, stroke_width, opacity, _dash = _style(primitive)
        shape = None
        if primitive.kind == "face" and len(primitive.points) >= 3:
            coordinates = [[x * MM_TO_EMU, y * MM_TO_EMU] for x, y in primitive.points]
            builder = slide.shapes.build_freeform(coordinates[0][0], coordinates[0][1])
            builder.add_line_segments([(x, y) for x, y in coordinates[1:]], close=True)
            shape = builder.convert_to_shape()
        elif primitive.kind in {"edge", "polyline", "leader", "arrow", "bezier"}:
            route_points = _bezier_samples(primitive.points, 25) if primitive.kind == "bezier" else primitive.points
            for segment_index, (start, end) in enumerate(zip(route_points, route_points[1:])):
                connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Mm(start[0]), Mm(start[1]), Mm(end[0]), Mm(end[1]))
                connector.name = f"NNDV editable {primitive.kind} {primitive.object_id} {segment_index:03d}"
                connector._element.nvCxnSpPr.cNvPr.set(
                    "descr",
                    json.dumps(
                        _presentation_binding(primitive, projected),
                        sort_keys=True,
                        separators=(",", ":"),
                    ),
                )
                connector.line.color.rgb = _pptx_colour(stroke)
                connector.line.width = Pt(stroke_width)
                if primitive.kind == "arrow" and segment_index == len(route_points) - 2:
                    line_xml = connector.line._get_or_add_ln()
                    tail = OxmlElement("a:tailEnd")
                    tail.set("type", "triangle")
                    line_xml.append(tail)
            continue
        elif primitive.kind == "label" and primitive.points:
            x, y = primitive.points[0]
            bbox = primitive.style.get("text_bbox_mm", [x, y, x + 30.0, y + 5.0])
            shape = slide.shapes.add_textbox(Mm(float(bbox[0])), Mm(float(bbox[1])), Mm(float(bbox[2]) - float(bbox[0])), Mm(float(bbox[3]) - float(bbox[1])))
            frame = shape.text_frame
            frame.clear()
            frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Emu(0)
            frame.vertical_anchor = MSO_ANCHOR.TOP
            paragraph = frame.paragraphs[0]
            paragraph.text = primitive.text
            paragraph.font.name = str(primitive.style.get("font_family", projected.options.font_family)).split(",")[0]
            paragraph.font.size = Pt(max(7.0, float(primitive.style.get("font_size_pt", 7.0))))
            paragraph.font.color.rgb = _pptx_colour(fill)
        if shape is not None:
            shape.name = f"NNDV editable {primitive.kind} {primitive.object_id} {index:04d}"
            shape._element.nvSpPr.cNvPr.set(
                "descr",
                json.dumps(
                    _presentation_binding(primitive, projected),
                    sort_keys=True,
                    separators=(",", ":"),
                ),
            )
            if primitive.kind == "face":
                shape.fill.solid()
                shape.fill.fore_color.rgb = _pptx_colour(fill)
                shape.fill.transparency = round((1.0 - opacity) * 100)
                shape.line.color.rgb = _pptx_colour(stroke)
                shape.line.width = Pt(stroke_width)
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            slide_xml = archive.read("ppt/slides/slide1.xml")
    except (OSError, KeyError, zipfile.BadZipFile) as exc:
        raise ExportError("Editable Scene PPTX is not a valid presentation package") from exc
    if "ppt/presentation.xml" not in names or (b"NNDV editable" not in slide_xml and b"NNDV Scene background" not in slide_xml):
        raise ExportError("Editable Scene PPTX contains no editable vector shapes")
    if any(name.startswith("ppt/media/") for name in names):
        raise ExportError("Editable Scene PPTX unexpectedly contains flattened media")
    return path


def _safe_script_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).replace("&", "\\u0026").replace("<", "\\u003c").replace(">", "\\u003e")


def render_scene_html(
    scene_or_projection: Any,
    *,
    camera_id: str | None = None,
    options: ProjectionOptions | Mapping[str, Any] | None = None,
) -> str:
    """Return a standalone, dependency-free pan/zoom/select SVG viewer."""

    projected = _projection(scene_or_projection, camera_id=camera_id, options=options)
    svg = render_scene_svg(projected)
    if svg.startswith("<?xml"):
        svg = svg.split("\n", 1)[1]
    projection_json = _safe_script_json(projected.to_dict())
    payload = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; script-src 'unsafe-inline'; img-src data:">
<title>NN_DaVinci Scene {_esc(projected.scene_id)}</title>
<style>
:root{{color-scheme:light dark;font-family:system-ui,-apple-system,"Segoe UI",sans-serif}}body{{margin:0;background:#e2e8f0;color:#0f172a}}
header{{height:48px;display:flex;align-items:center;gap:8px;padding:0 14px;background:#fff;border-bottom:1px solid #cbd5e1}}
button{{height:32px;border:1px solid #94a3b8;border-radius:6px;background:#fff;color:#0f172a;padding:0 10px}}
#stage{{height:calc(100vh - 49px);overflow:hidden;touch-action:none;background:#cbd5e1}}#stage svg{{width:100%;height:100%;display:block}}
[data-object-id]{{cursor:pointer}}[data-object-id].selected{{filter:drop-shadow(0 0 1.2mm #2563eb)}}#status{{margin-left:auto;font-size:12px;color:#475569}}
.labels-hidden [data-kind="label"],.labels-hidden [data-kind="label-halo"],.labels-hidden [data-kind="leader"]{{display:none}}
</style>
</head>
<body>
<header><strong>Scene vector view</strong><button id="reset" type="button">Reset view</button><button id="labels" type="button">Labels</button><span id="status">Camera: {_esc(projected.camera_id)}</span></header>
<main id="stage" aria-label="Interactive offline Scene projection">{svg}</main>
<script id="nndv-scene-projection" type="application/json">{projection_json}</script>
<script>
(()=>{{'use strict';const stage=document.getElementById('stage'),svg=stage.querySelector('svg'),status=document.getElementById('status');
const original=[0,0,{projected.width_mm},{projected.height_mm}],box=[...original];let drag=null;
const apply=()=>svg.setAttribute('viewBox',box.join(' '));document.getElementById('reset').onclick=()=>{{box.splice(0,4,...original);apply()}};
document.getElementById('labels').onclick=()=>stage.classList.toggle('labels-hidden');
stage.addEventListener('wheel',event=>{{event.preventDefault();const ratio=Math.exp(event.deltaY*.001);const rect=stage.getBoundingClientRect();const px=box[0]+event.offsetX/rect.width*box[2],py=box[1]+event.offsetY/rect.height*box[3];box[0]=px-(px-box[0])*ratio;box[1]=py-(py-box[1])*ratio;box[2]*=ratio;box[3]*=ratio;apply()}},{{passive:false}});
stage.addEventListener('pointerdown',event=>{{drag=[event.clientX,event.clientY,...box];stage.setPointerCapture(event.pointerId)}});
stage.addEventListener('pointermove',event=>{{if(!drag)return;const rect=stage.getBoundingClientRect();box[0]=drag[2]-(event.clientX-drag[0])/rect.width*drag[4];box[1]=drag[3]-(event.clientY-drag[1])/rect.height*drag[5];apply()}});
stage.addEventListener('pointerup',()=>drag=null);stage.addEventListener('click',event=>{{const item=event.target.closest('[data-object-id]');if(!item)return;stage.querySelectorAll('.selected').forEach(node=>node.classList.remove('selected'));const id=item.dataset.objectId;stage.querySelectorAll('[data-object-id="'+CSS.escape(id)+'"]').forEach(node=>node.classList.add('selected'));status.textContent='Selected: '+id}});apply();
}})();
</script>
</body>
</html>
"""
    validate_scene_html(payload)
    return payload


def validate_scene_html(payload: str) -> None:
    lowered = payload.lower()
    if "<script" not in lowered or "nndv-scene-projection" not in lowered or "<svg" not in lowered:
        raise ExportError("Offline Scene HTML is missing its embedded projection/viewer")
    prohibited_patterns = (
        r"<(?:script|img|link)[^>]+(?:src|href)\s*=\s*['\"]\s*(?:https?:)?//",
        r"\b(?:fetch|xmlhttprequest|websocket)\s*\(",
        r"javascript:",
    )
    if any(re.search(pattern, lowered) for pattern in prohibited_patterns):
        raise ExportError("Offline Scene HTML contains a network or executable dependency")


def export_scene_pdf(
    scene_or_projection: Any,
    destination: str | Path,
    *,
    camera_id: str | None = None,
    options: ProjectionOptions | Mapping[str, Any] | None = None,
) -> Path:
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(render_scene_pdf(scene_or_projection, camera_id=camera_id, options=options))
    validate_scene_pdf(path.read_bytes())
    return path


def export_scene_eps(
    scene_or_projection: Any,
    destination: str | Path,
    *,
    camera_id: str | None = None,
    options: ProjectionOptions | Mapping[str, Any] | None = None,
) -> Path:
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(render_scene_eps(scene_or_projection, camera_id=camera_id, options=options))
    validate_scene_eps(path.read_bytes())
    return path


def export_scene_html(
    scene_or_projection: Any,
    destination: str | Path,
    *,
    camera_id: str | None = None,
    options: ProjectionOptions | Mapping[str, Any] | None = None,
) -> Path:
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(render_scene_html(scene_or_projection, camera_id=camera_id, options=options), encoding="utf-8")
    validate_scene_html(path.read_text(encoding="utf-8"))
    return path


def scene_export_policy(
    scene_or_projection: Any, *, camera_id: str | None = None, options: ProjectionOptions | Mapping[str, Any] | None = None
) -> dict[str, Any]:
    projected = _projection(scene_or_projection, camera_id=camera_id, options=options)
    return {
        "schema_version": SCENE_EXPORT_VERSION,
        "source_digest": projected.source_digest,
        "camera_id": projected.camera_id,
        "physical_page_mm": [projected.width_mm, projected.height_mm],
        "stroke_width_pt": projected.options.stroke_width_pt,
        "minimum_label_pt": projected.options.min_label_pt,
        "hidden_line_removal": projected.options.hidden_edges,
        "backface_culling": projected.options.backface_culling,
        "formats": {
            "svg": "native polygons, polylines, cubic paths and text",
            "pdf": "native PDF paths and text; no Image XObjects",
            "tikz": "native TikZ paths, cubic controls and nodes",
            "pptx": "editable DrawingML shapes and connectors; no flattened media",
            "png": f"direct CPU raster proof at {PNG_EXPORT_DPI} DPI",
            "eps": "native EPSF Level 3 PostScript paths and text; no image operators",
            "html": "standalone offline SVG pan/zoom/select viewer",
            "json": "source-preserving Scene IR JSON",
            "gltf": "glTF 2.0 with embedded binary buffer",
            "glb": "binary glTF 2.0",
        },
    }


def export_scene(
    scene: Any,
    output: str | Path,
    *,
    formats: Iterable[str] = ("svg",),
    camera_id: str | None = None,
    options: ProjectionOptions | Mapping[str, Any] | None = None,
) -> list[Path]:
    """Export Scene IR to requested 2D/3D formats in request order."""

    validate_scene_assets(scene)
    stem = Path(output).with_suffix("")
    stem.parent.mkdir(parents=True, exist_ok=True)
    # Project once so every 2D format shares byte-for-byte geometry decisions.
    projected: ProjectedScene | None = None

    def projection() -> ProjectedScene:
        nonlocal projected
        if projected is None:
            publication_options = options if options is not None else ProjectionOptions(
                auto_frame=True,
                density="paper",
                target_occupancy=0.72,
            )
            projected = project_scene(scene, camera_id=camera_id, options=publication_options)
        return projected

    outputs: list[Path] = []
    for requested in formats:
        name = str(requested).lower().lstrip(".")
        if name == "svg":
            destination = stem.with_suffix(".svg")
            destination.write_text(render_scene_svg(projection()), encoding="utf-8")
            validate_scene_svg(destination.read_text(encoding="utf-8"))
        elif name == "pdf":
            destination = export_scene_pdf(projection(), stem.with_suffix(".pdf"))
        elif name in {"tikz", "tex"}:
            destination = stem.with_suffix(".tex")
            destination.write_text(render_scene_tikz(projection()), encoding="utf-8")
        elif name in {"pptx", "powerpoint"}:
            destination = export_scene_pptx(projection(), stem.with_suffix(".pptx"))
        elif name == "png":
            destination = export_scene_png(projection(), stem.with_suffix(".png"))
        elif name in {"eps", "epsf"}:
            destination = export_scene_eps(projection(), stem.with_suffix(".eps"))
        elif name in {"html", "htm"}:
            destination = export_scene_html(projection(), stem.with_suffix(".html"))
        elif name in {"json", "scene-json", "scene_json"}:
            destination = export_scene_json(scene, stem.parent / f"{stem.name}.scene.json")
        elif name == "gltf":
            destination = export_scene_gltf(scene, stem.with_suffix(".gltf"))
        elif name == "glb":
            destination = export_scene_glb(scene, stem.with_suffix(".glb"))
        else:
            raise ExportError("Scene IR supports SVG, PDF, TikZ/TeX, editable PPTX, 300-DPI PNG, vector EPS, offline HTML, Scene JSON, glTF and GLB")
        if not destination.is_file() or destination.stat().st_size <= 0:
            raise ExportError(f"Scene {name} exporter did not produce a non-empty output")
        outputs.append(destination)
    return outputs


__all__ = [
    "PNG_EXPORT_DPI",
    "SCENE_EXPORT_VERSION",
    "export_scene",
    "export_scene_eps",
    "export_scene_html",
    "export_scene_pdf",
    "export_scene_png",
    "export_scene_pptx",
    "render_scene_eps",
    "render_scene_html",
    "render_scene_pdf",
    "render_scene_svg",
    "render_scene_tikz",
    "scene_export_policy",
    "validate_scene_eps",
    "validate_scene_html",
    "validate_scene_pdf",
    "validate_scene_svg",
]
