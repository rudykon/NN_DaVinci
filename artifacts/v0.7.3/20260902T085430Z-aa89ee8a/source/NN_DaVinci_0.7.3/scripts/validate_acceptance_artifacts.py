#!/usr/bin/env python3
"""Perform parser/signature/editability/security checks on fresh exports."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def check(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def validate(directory: Path) -> dict[str, int]:
    stem = directory / "resnet-v02"
    paths = {suffix: stem.with_suffix(suffix) for suffix in (".svg", ".pdf", ".tex", ".png", ".eps", ".pptx", ".html")}
    for suffix, path in paths.items():
        check(path.is_file() and path.stat().st_size > 32, f"missing/empty {suffix}: {path}")
    check(paths[".pdf"].read_bytes().startswith(b"%PDF-"), "invalid PDF signature")
    check(paths[".png"].read_bytes().startswith(b"\x89PNG\r\n\x1a\n"), "invalid PNG signature")
    check(paths[".eps"].read_bytes().startswith(b"%!PS-Adobe"), "invalid EPS signature")
    check(zipfile.is_zipfile(paths[".pptx"]), "PPTX is not an OOXML ZIP")
    check(paths[".tex"].read_text(encoding="utf-8").lstrip().startswith("\\documentclass"), "TikZ is not standalone")

    root = ElementTree.parse(paths[".svg"]).getroot()
    check(root.tag.endswith("svg") and len(root.findall(".//{http://www.w3.org/2000/svg}g")) > 0, "SVG parser found no vector groups")

    presentation = Presentation(paths[".pptx"])
    shapes = [shape for slide in presentation.slides for shape in slide.shapes]
    pictures = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    connectors = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE]
    text_shapes = [shape for shape in shapes if getattr(shape, "has_text_frame", False)]
    check(not pictures, "PPTX contains a flattened picture")
    check(len(text_shapes) >= 5, "PPTX lacks editable node shapes")
    check(len(connectors) >= 5, "PPTX lacks editable connectors")

    html = paths[".html"].read_text(encoding="utf-8")
    check("Content-Security-Policy" in html and "default-src 'none'" in html, "HTML lacks restrictive CSP")
    check("replaceChildren" in html and ".innerHTML" not in html, "HTML inspector is not DOM-XSS safe")
    check(not re.search(r"<(?:script|link|img)[^>]+(?:src|href)=[\"']https?://", html, re.I), "HTML has an external resource")
    check("<svg" in html and "graph-data" in html, "HTML is not self-contained")
    return {"checks": 20, "failures": 0, "errors": 0, "skipped": 0}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    summary = validate(args.directory)
    if args.output:
        args.output.write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(f"NNDV_ARTIFACT_SUMMARY={json.dumps(summary, sort_keys=True)}")


if __name__ == "__main__":
    main()
