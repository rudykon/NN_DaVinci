#!/usr/bin/env python3
"""Reject editor affordances in publication outputs while preserving MoE provenance."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import subprocess

from pptx import Presentation


FORBIDDEN = ("⊞", "editor-control", "resize-handle", "selection-box")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("authority_directory", type=Path)
    parser.add_argument("acceptance_directory", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    failures: list[str] = []
    checked: list[str] = []
    for path in sorted(args.authority_directory.glob("*.svg")):
        text = path.read_text(encoding="utf-8")
        checked.append(str(path))
        if any(token in text for token in FORBIDDEN):
            failures.append(f"{path.name}: publication SVG contains editor affordance")
    moe = (args.authority_directory / "moe.svg").read_text(encoding="utf-8")
    if "Expert ×4" not in moe:
        failures.append("moe.svg: Expert ×4 provenance is absent")
    for path in sorted(args.authority_directory.glob("*.pdf")):
        extracted = subprocess.run(["pdftotext", str(path), "-"], check=True, capture_output=True, text=True).stdout
        checked.append(str(path))
        if any(token in extracted for token in FORBIDDEN):
            failures.append(f"{path.name}: publication PDF contains editor affordance")
    for path in sorted(args.authority_directory.glob("*.tex")):
        text = path.read_text(encoding="utf-8")
        checked.append(str(path))
        if any(token in text for token in FORBIDDEN):
            failures.append(f"{path.name}: publication TikZ contains editor affordance")
    pptx_path = args.acceptance_directory / "resnet-v02.pptx"
    presentation = Presentation(pptx_path)
    pptx_text = "\n".join(
        shape.text for slide in presentation.slides for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    checked.append(str(pptx_path))
    if any(token in pptx_text for token in FORBIDDEN):
        failures.append("PPTX contains editor affordance")
    for suffix in (".png", ".eps"):
        path = args.acceptance_directory / f"resnet-v02{suffix}"
        checked.append(str(path))
        if not path.is_file() or path.stat().st_size == 0:
            failures.append(f"publication {suffix} is absent")
    report = {"checked": checked, "moe_provenance": "Expert ×4", "failures": failures, "passed": not failures}
    args.output.write_text(json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps({"publication_outputs": len(checked), "passed": not failures}, sort_keys=True))
    if failures:
        raise SystemExit("\n".join(failures))


if __name__ == "__main__":
    main()
