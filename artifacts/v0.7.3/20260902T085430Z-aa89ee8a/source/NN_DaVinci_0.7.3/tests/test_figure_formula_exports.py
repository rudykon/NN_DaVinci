from __future__ import annotations

import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

from PIL import Image
from pptx import Presentation

from nn_davinci.figure_export import compile_figure, export_figure, render_figure_svg, render_figure_tikz
from nn_davinci.figure_ir import FigureObject, FigureProvenance, FigureStyle, new_figure
from nn_davinci.formula import FORMULA_SYNTAX_VERSION, normalize_formula


FORMULAS = (
    r"\alpha_i^2 + \beta \leq \sqrt{x} + \frac{1}{N}",
    r"\sum_i q_i \times k_i \to attention",
    r"A = \begin{bmatrix}a & b \\ c & \delta\end{bmatrix}",
    r"\input{/etc/passwd}",
)


def _formula_figure():
    figure = new_figure("Editable scientific formulas")
    panel = next(figure.iter_panels())
    layer = next(item for item in panel.layers if item.role == "author-annotation")
    geometries = (
        {"x": 12.0, "y": 14.0, "width": 150.0, "height": 14.0},
        {"x": 12.0, "y": 34.0, "width": 150.0, "height": 14.0},
        {"x": 12.0, "y": 54.0, "width": 150.0, "height": 22.0},
        {"x": 12.0, "y": 83.0, "width": 150.0, "height": 18.0},
    )
    for index, (source, geometry) in enumerate(zip(FORMULAS, geometries, strict=True)):
        layer.objects.append(FigureObject.create(
            "equation",
            f"Equation {index + 1}",
            geometry,
            FigureProvenance.author("Formula export regression fixture."),
            style=FigureStyle(overrides={"font_size": 8.0}),
            metadata={"text": source, "anchor": "start", "vertical_anchor": "middle"},
            identity=f"formula-export:{index}",
            order=index,
        ))
    return figure.validate()


class FigureFormulaCompilationTests(unittest.TestCase):
    def test_compiled_primitives_retain_complete_normalization_record(self) -> None:
        figure = _formula_figure()
        primitives = [
            primitive
            for _page_id, _panel_id, _layer_id, primitive in compile_figure(figure)
            if primitive.values.get("role") == "equation"
        ]
        self.assertEqual(4, len(primitives))
        for primitive, source in zip(primitives, FORMULAS, strict=True):
            normalized = normalize_formula(source)
            self.assertEqual(normalized.text, primitive.values["text"])
            self.assertEqual(source, primitive.values["formula_source"])
            self.assertEqual(FORMULA_SYNTAX_VERSION, primitive.values["formula_syntax_version"])
            self.assertEqual(normalized.valid, primitive.values["formula_valid"])
            self.assertEqual(normalized.fallback, primitive.values["formula_fallback"])
            self.assertEqual(normalized.error, primitive.values["formula_error"])
            self.assertEqual(list(normalized.features), primitive.values["formula_features"])
        self.assertIn("αᵢ²", primitives[0].values["text"])
        self.assertIn("β ≤ √(x)", primitives[0].values["text"])
        self.assertIn("∑ᵢ", primitives[1].values["text"])
        self.assertEqual(2, len(primitives[2].values["lines"]))
        self.assertFalse(primitives[3].values["formula_valid"])
        self.assertIn(r"Invalid formula: \input{/etc/passwd}", primitives[3].values["text"])

    def test_svg_serializes_safe_formula_state_as_editable_text(self) -> None:
        svg = render_figure_svg(_formula_figure())
        root = ET.fromstring(svg)
        namespace = {"svg": "http://www.w3.org/2000/svg"}
        equations = root.findall('.//svg:text[@data-text-role="equation"]', namespace)
        self.assertEqual(4, len(equations))
        for element, source in zip(equations, FORMULAS, strict=True):
            normalized = normalize_formula(source)
            self.assertEqual(source, element.attrib["data-formula-source"])
            self.assertEqual(FORMULA_SYNTAX_VERSION, element.attrib["data-formula-syntax-version"])
            self.assertEqual(str(normalized.valid).lower(), element.attrib["data-formula-valid"])
            self.assertEqual(str(normalized.fallback).lower(), element.attrib["data-formula-fallback"])
            self.assertEqual(normalized.error, element.attrib["data-formula-error"])
            self.assertEqual(list(normalized.features), json.loads(element.attrib["data-formula-features"]))
            self.assertEqual("editable-text", element.attrib["data-formula-rendering"])
            self.assertEqual(normalized.text.replace("\n", ""), "".join(element.itertext()))
        self.assertFalse(root.findall('.//svg:image[@data-text-role="equation"]', namespace))

    def test_tikz_uses_safe_editable_commands_and_never_executes_invalid_source(self) -> None:
        tikz = render_figure_tikz(_formula_figure())
        self.assertIn(r"\ensuremath{\alpha}\textsubscript{i}\textsuperscript{2}", tikz)
        self.assertIn(r"\ensuremath{\beta} \ensuremath{\leq} \ensuremath{\surd}(x)", tikz)
        self.assertIn(r"\ensuremath{\sum}\textsubscript{i}", tikz)
        self.assertIn(r"[a b]", tikz)
        self.assertIn(r"[c \ensuremath{\delta}]", tikz)
        self.assertNotIn(r"\input{/etc/passwd}", tikz)
        self.assertIn(r"Invalid formula: \textbackslash{}input\{/etc/passwd\}", tikz)


@unittest.skipUnless(shutil.which("pdflatex") and shutil.which("pdftops"), "TeX and Poppler are required")
class FigureFormulaCrossFormatTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temporary = tempfile.TemporaryDirectory(prefix="nndv-formula-export-test-")
        cls.root = Path(cls.temporary.name)
        cls.figure = _formula_figure()
        cls.outputs = export_figure(
            cls.figure,
            cls.root / "formulas",
            formats=("svg", "pdf", "tikz", "pptx", "png", "eps", "html"),
        )

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temporary.cleanup()

    def test_all_seven_formats_derive_from_one_editable_formula_text_model(self) -> None:
        self.assertEqual(
            [".svg", ".pdf", ".tex", ".pptx", ".png", ".eps", ".html"],
            [path.suffix for path in self.outputs],
        )
        svg = (self.root / "formulas.svg").read_text(encoding="utf-8")
        html = (self.root / "formulas.html").read_text(encoding="utf-8")
        self.assertIn("αᵢ²", svg)
        self.assertIn("αᵢ²", html)
        self.assertIn('data-formula-rendering="editable-text"', svg)
        self.assertIn('data-formula-rendering="editable-text"', html)
        self.assertTrue((self.root / "formulas.pdf").read_bytes().startswith(b"%PDF"))
        with Image.open(self.root / "formulas.png") as image:
            self.assertEqual("RGBA", image.mode)
            self.assertEqual(self.figure.digest(), json.loads(image.info["nndv.figure_export"])["source_figure_digest"])
        eps = (self.root / "formulas.eps").read_bytes()
        self.assertIn(b"%%NNDVVectorPolicy: SVG-to-PDF-to-Poppler-EPS", eps)
        self.assertIn(b"%%BeginResource: font", eps)

        presentation = Presentation(self.root / "formulas.pptx")
        text = "\n".join(
            shape.text
            for shape in presentation.slides[0].shapes
            if getattr(shape, "has_text_frame", False)
        )
        self.assertIn("αᵢ²", text)
        self.assertIn("[a b]", text)
        self.assertIn(r"Invalid formula: \input{/etc/passwd}", text)
        with zipfile.ZipFile(self.root / "formulas.pptx") as archive:
            slide = archive.read("ppt/slides/slide1.xml").decode("utf-8")
            self.assertIn("NNDV editable formula", slide)
            self.assertIn("formula_syntax_version", slide)
            self.assertNotIn("&lt;p:pic", slide)

    def test_generated_tikz_compiles_without_shell_escape(self) -> None:
        command = [
            "pdflatex", "-no-shell-escape", "-interaction=nonstopmode", "-halt-on-error",
            "-output-directory", str(self.root), str(self.root / "formulas.tex"),
        ]
        result = subprocess.run(command, check=False, capture_output=True, text=True)
        self.assertEqual(0, result.returncode, result.stdout + result.stderr)
        self.assertTrue((self.root / "formulas.pdf").read_bytes().startswith(b"%PDF"))


if __name__ == "__main__":
    unittest.main()
