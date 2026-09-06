from __future__ import annotations

import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from PIL import Image
from pptx import Presentation

from nn_davinci.errors import ExportError, OptionalDependencyError
from nn_davinci.figure_export import (
    PNG_EXPORT_DPI,
    export_figure,
    export_figure_eps_pages,
    export_figure_pdf,
    figure_export_policy,
    render_figure_html,
    render_figure_svg_pages,
)
from nn_davinci.figure_ir import FigureObject, FigureProvenance, new_figure


def _add_content(figure, label: str) -> None:
    panel = next(figure.iter_panels())
    layer = next(item for item in panel.layers if item.role == "model-data")
    layer.objects.append(FigureObject.create(
        "node-glyph",
        label,
        {"x": 10.0, "y": 25.0, "width": 50.0, "height": 15.0},
        FigureProvenance.author("Multi-page export test content."),
        identity=f"multi-page:{label}",
    ))


def _heterogeneous_figure():
    early = new_figure("Early portrait page", page_preset="single-column")
    late = new_figure("Late landscape page", page_preset="double-column")
    _add_content(early, "EARLY_PAGE_CONTENT")
    _add_content(late, "LATE_PAGE_CONTENT")
    early.pages[0].order = 10
    late.pages[0].order = 20
    # Deliberately store pages out of order; every exporter must use page order.
    late.pages = [late.pages[0], early.pages[0]]
    return late.validate()


@unittest.skipUnless(shutil.which("pdfunite") and shutil.which("pdftops"), "Poppler export tools are required")
class FigureMultiPageExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temporary = tempfile.TemporaryDirectory(prefix="nndv-multi-page-test-")
        cls.root = Path(cls.temporary.name)
        cls.figure = _heterogeneous_figure()
        cls.outputs = export_figure(
            cls.figure,
            cls.root / "figure",
            formats=("svg", "pdf", "tikz", "pptx", "png", "eps", "html"),
        )

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temporary.cleanup()

    def test_page_first_counts_signatures_and_order(self) -> None:
        self.assertEqual(11, len(self.outputs))
        self.assertEqual(
            [
                "figure.page-001.svg", "figure.page-002.svg", "figure.pdf",
                "figure.page-001.tex", "figure.page-002.tex", "figure.pptx",
                "figure.page-001.png", "figure.page-002.png",
                "figure.page-001.eps", "figure.page-002.eps", "figure.html",
            ],
            [path.name for path in self.outputs],
        )
        signatures = {
            ".svg": b"<?xml", ".pdf": b"%PDF", ".tex": b"\\documentclass",
            ".pptx": b"PK", ".png": b"\x89PNG\r\n\x1a\n",
            ".eps": b"%!PS-Adobe-3.0 EPSF-3.0", ".html": b"<!doctype html>",
        }
        for path in self.outputs:
            with self.subTest(path=path.name):
                self.assertTrue(path.read_bytes().startswith(signatures[path.suffix]))
        first_svg = (self.root / "figure.page-001.svg").read_text(encoding="utf-8")
        second_svg = (self.root / "figure.page-002.svg").read_text(encoding="utf-8")
        self.assertIn("EARLY_PAGE_CONTENT", first_svg)
        self.assertNotIn("LATE_PAGE_CONTENT", first_svg)
        self.assertIn("LATE_PAGE_CONTENT", second_svg)
        self.assertNotIn("EARLY_PAGE_CONTENT", second_svg)

    def test_pdf_is_paginated_and_preserves_heterogeneous_page_sizes(self) -> None:
        result = subprocess.run(
            ["pdfinfo", "-f", "1", "-l", "2", str(self.root / "figure.pdf")],
            check=True,
            capture_output=True,
            text=True,
        )
        self.assertIn("Pages:           2", result.stdout)
        self.assertRegex(result.stdout, r"Page\s+1 size:\s+249\.44\d x 334\.48\d pts")
        self.assertRegex(result.stdout, r"Page\s+2 size:\s+504\.56\d x 334\.48\d pts")

    def test_pptx_has_one_ordered_editable_slide_per_page(self) -> None:
        presentation = Presentation(self.root / "figure.pptx")
        self.assertEqual(2, len(presentation.slides))
        self.assertAlmostEqual(178.0, presentation.slide_width / 36_000, places=3)
        self.assertAlmostEqual(118.0, presentation.slide_height / 36_000, places=3)
        expected = (("EARLY_PAGE_CONTENT", self.figure.pages[1].id), ("LATE_PAGE_CONTENT", self.figure.pages[0].id))
        for index, (label, page_id) in enumerate(expected):
            shapes = list(presentation.slides[index].shapes)
            self.assertTrue(shapes[0].name.startswith(f"NNDV Page {index + 1:03d} {page_id}"))
            self.assertIn(label, [shape.text for shape in shapes if getattr(shape, "has_text_frame", False)])
        with zipfile.ZipFile(self.root / "figure.pptx") as archive:
            for slide_number in (1, 2):
                xml = archive.read(f"ppt/slides/slide{slide_number}.xml")
                self.assertIn(b"NNDV editable", xml)
                self.assertNotIn(b"<p:pic>", xml)

    def test_png_is_300_dpi_rgba_with_explicit_alpha_metadata(self) -> None:
        expected_pixels = ((1039, 1394), (2102, 1394))
        for index, expected_size in enumerate(expected_pixels, start=1):
            with Image.open(self.root / f"figure.page-{index:03d}.png") as image:
                self.assertEqual("RGBA", image.mode)
                self.assertEqual(expected_size, image.size)
                self.assertAlmostEqual(PNG_EXPORT_DPI, image.info["dpi"][0], places=2)
                metadata = json.loads(image.info["nndv.figure_export"])
                self.assertEqual(index, metadata["page_number"])
                self.assertEqual(PNG_EXPORT_DPI, metadata["dpi"])
                self.assertIn("straight-alpha RGBA", metadata["alpha_semantics"])

    def test_eps_is_vector_level_three_and_embeds_fonts(self) -> None:
        for index in (1, 2):
            payload = (self.root / f"figure.page-{index:03d}.eps").read_bytes()
            self.assertIn(b"%%NNDVVectorPolicy: SVG-to-PDF-to-Poppler-EPS", payload)
            self.assertIn(b"%%NNDVFontPolicy: embedded subset resources required", payload)
            self.assertIn(b"%%BeginResource: font", payload)

    def test_html_is_offline_editable_and_preserves_complete_source(self) -> None:
        document = (self.root / "figure.html").read_text(encoding="utf-8")
        self.assertIn("default-src 'none'", document)
        self.assertIn("Download edited page SVG", document)
        self.assertIn("data-page-index=\"0\"", document)
        self.assertIn("data-page-index=\"1\"", document)
        self.assertIn("EARLY_PAGE_CONTENT", document)
        self.assertIn("LATE_PAGE_CONTENT", document)
        self.assertEqual(1, document.count('id="nndv-arrow-page-1"'))
        self.assertEqual(1, document.count('id="nndv-arrow-page-2"'))
        self.assertNotIn('id="nndv-arrow"', document)
        match = re.search(
            r'<script id="nndv-figure-ir" type="application/json">(.*?)</script>',
            document,
            flags=re.DOTALL,
        )
        self.assertIsNotNone(match)
        source = json.loads(match.group(1))
        self.assertEqual(self.figure.to_dict(), source)

    def test_policy_makes_every_heterogeneous_page_decision_explicit(self) -> None:
        policy = figure_export_policy(self.figure)
        self.assertEqual([self.figure.pages[1].id, self.figure.pages[0].id], policy["page_order"])
        self.assertEqual(["portrait", "landscape"], [page["orientation"] for page in policy["pages"]])
        self.assertEqual("preserved", policy["formats"]["pdf"]["heterogeneous_page_sizes"])
        self.assertIn("centered at 1:1", policy["formats"]["pptx"]["heterogeneous_page_sizes"])
        self.assertEqual(0, policy["formats"]["html"]["network_dependencies"])


class FigureExportFailureAndCompatibilityTests(unittest.TestCase):
    def test_single_page_keeps_legacy_stem_for_all_seven_formats(self) -> None:
        figure = new_figure("Single-page compatibility")
        _add_content(figure, "SINGLE_PAGE_CONTENT")
        if not (shutil.which("pdftops") and shutil.which("pdfunite")):
            self.skipTest("Poppler export tools are required")
        with tempfile.TemporaryDirectory(prefix="nndv-single-page-test-") as directory:
            outputs = export_figure(
                figure,
                Path(directory) / "figure",
                formats=("svg", "pdf", "tikz", "pptx", "png", "eps", "html"),
            )
            self.assertEqual(
                ["figure.svg", "figure.pdf", "figure.tex", "figure.pptx", "figure.png", "figure.eps", "figure.html"],
                [path.name for path in outputs],
            )

    def test_missing_page_converters_fail_actionably_without_raster_fallback(self) -> None:
        figure = _heterogeneous_figure()
        with tempfile.TemporaryDirectory(prefix="nndv-missing-converter-test-") as directory:
            with patch("nn_davinci.figure_export.shutil.which", return_value=None):
                with self.assertRaisesRegex(OptionalDependencyError, "pdfunite"):
                    export_figure_pdf(figure, Path(directory) / "figure.pdf")
                with self.assertRaisesRegex(OptionalDependencyError, "will not silently rasterize"):
                    export_figure_eps_pages(figure, Path(directory) / "figure")

    def test_offline_html_rejects_network_image_dependencies(self) -> None:
        figure = new_figure("External asset")
        panel = next(figure.iter_panels())
        layer = next(item for item in panel.layers if item.role == "author-annotation")
        layer.objects.append(FigureObject.create(
            "image", "remote",
            {"x": 10, "y": 20, "width": 20, "height": 20},
            FigureProvenance.author("Explicit test image."),
            metadata={"href": "https://example.invalid/image.png"},
        ))
        with self.assertRaisesRegex(ExportError, "embed it as a data URI"):
            render_figure_html(figure)

    def test_page_svg_projection_links_back_to_source_digest(self) -> None:
        figure = _heterogeneous_figure()
        pages = render_figure_svg_pages(figure)
        self.assertEqual(2, len(pages))
        for _page_id, svg in pages:
            self.assertIn(figure.digest(), svg)


if __name__ == "__main__":
    unittest.main()
