"""Deterministic vector rendering, round-trip import, and submission bundles."""

from __future__ import annotations

import base64
import copy
from dataclasses import dataclass
import html
import io
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
from typing import Any, Iterable
import xml.etree.ElementTree as ET

from .errors import ExportError, OptionalDependencyError, ValidationError
from .formula import normalize_formula
from .figure_ir import (
    FIGURE_IR_VERSION,
    FigureIR,
    FigureObject,
    FigureProvenance,
    new_figure,
)
from .scene_ir import empty_scene
from .tensor_geometry import TensorGeometrySpec, contrast_ratio, operator_symbol, relative_luminance, tensor_geometry
from .text_layout import TextLayoutError, abbreviated_label, layout_text_box, measure_text_line
from .units import pt_to_mm
from .version import __version__


FIGURE_SVG_METADATA_VERSION = "nndv-figure-svg-metadata-1"
FIGURE_EXPORT_VERSION = "1.1"
MM_TO_EMU = 36_000
PNG_EXPORT_DPI = 300
PAGE_FILE_TEMPLATE = ".page-{index:03d}"


def _esc(value: Any) -> str:
    return html.escape(str(value), quote=True)


def _hex(value: Any, fallback: str = "#334155") -> str:
    text = str(value or fallback)
    if len(text) == 7 and text.startswith("#") and all(ch in "0123456789abcdefABCDEF" for ch in text[1:]):
        return text.lower()
    return fallback


def _ordered_pages(figure: FigureIR) -> list[Any]:
    """Return pages in the one canonical export order."""

    return sorted(figure.pages, key=lambda item: (item.order, item.id))


def _page_figure(figure: FigureIR, page_index: int) -> FigureIR:
    """Create a validated, source-linked one-page projection of ``figure``."""

    pages = _ordered_pages(figure)
    if not 0 <= page_index < len(pages):
        raise ExportError(f"Figure page index {page_index} is outside 0..{len(pages) - 1}")
    source_digest = figure.digest()
    page = pages[page_index]
    document = figure.to_dict()
    page_documents = sorted(document["pages"], key=lambda item: (int(item.get("order", 0)), str(item.get("id", ""))))
    document["pages"] = [copy.deepcopy(page_documents[page_index])]
    document["metadata"] = copy.deepcopy(document.get("metadata", {}))
    document["metadata"]["page_export"] = {
        "schema_version": "nndv-page-export-1",
        "source_figure_digest": source_digest,
        "source_page_count": len(pages),
        "page_index": page_index,
        "page_number": page_index + 1,
        "page_id": page.id,
    }
    return FigureIR.from_dict(document)


def _page_destination(stem: Path, suffix: str, page_index: int, page_count: int) -> Path:
    """Keep legacy names for one page and number every page otherwise."""

    if page_count == 1:
        return stem.with_suffix(suffix)
    return stem.parent / f"{stem.name}{PAGE_FILE_TEMPLATE.format(index=page_index + 1)}{suffix}"


def figure_export_policy(figure: FigureIR) -> dict[str, Any]:
    """Describe the deterministic page and fidelity behavior of all formats."""

    figure.validate()
    pages = _ordered_pages(figure)
    page_records = [
        {
            "index": index,
            "number": index + 1,
            "id": page.id,
            "title": page.title,
            "width_mm": page.width_mm,
            "height_mm": page.height_mm,
            "orientation": "landscape" if page.width_mm > page.height_mm else (
                "portrait" if page.height_mm > page.width_mm else "square"
            ),
        }
        for index, page in enumerate(pages)
    ]
    return {
        "schema_version": "nndv-figure-export-policy-1",
        "export_version": FIGURE_EXPORT_VERSION,
        "source_figure_digest": figure.digest(),
        "page_order": [page.id for page in pages],
        "pages": page_records,
        "formats": {
            "svg": {
                "multi_page": "one numbered, source-linked SVG per page",
                "single_page_compatibility": "<stem>.svg",
            },
            "pdf": {
                "multi_page": "one PDF page per FigurePage in canonical order",
                "heterogeneous_page_sizes": "preserved",
            },
            "tikz": {
                "multi_page": "one numbered standalone TeX source per page",
                "single_page_compatibility": "<stem>.tex",
            },
            "pptx": {
                "multi_page": "one editable slide per FigurePage in canonical order",
                "heterogeneous_page_sizes": (
                    "PowerPoint has one presentation-wide slide size; the canvas is the maximum page width and "
                    "height, and each page is centered at 1:1 physical scale without cropping"
                ),
                "canvas_width_mm": max(page.width_mm for page in pages),
                "canvas_height_mm": max(page.height_mm for page in pages),
            },
            "png": {
                "multi_page": "one numbered PNG per page",
                "dpi": PNG_EXPORT_DPI,
                "alpha_semantics": (
                    "straight-alpha RGBA; the authored FigureIR background is rendered exactly, and transparent "
                    "source assets retain alpha until composited"
                ),
            },
            "eps": {
                "multi_page": "one numbered EPSF Level 3 file per page",
                "conversion": "SVG vector -> PDF vector -> Poppler EPS; no page rasterization fallback",
                "fonts": "subset fonts are embedded; export fails if embedded font resources are absent",
            },
            "html": {
                "multi_page": "one offline viewer containing every page",
                "source_preservation": "the original canonical FigureIR and per-page SVG DOM remain embedded",
                "network_dependencies": 0,
            },
        },
    }


def _style(figure: FigureIR, item: FigureObject) -> dict[str, Any]:
    values = item.style.resolved(figure.design_tokens)
    stroke_width_pt = float(values.get("stroke_width", figure.design_tokens["line.width.pt"]))
    font_size_pt = float(values.get("font_size", figure.design_tokens["font.size.pt"]))
    radius_pt = float(values.get("radius", figure.design_tokens["corner.radius.pt"]))
    arrow_size_pt = float(values.get("arrow_size", figure.design_tokens["arrow.size.pt"]))
    if stroke_width_pt <= 0 or font_size_pt <= 0 or radius_pt < 0 or arrow_size_pt <= 0:
        raise ValidationError(f"Figure object {item.id!r} has a non-positive physical style value")
    return {
        "fill": _hex(values.get("fill", figure.design_tokens["color.operator"]), "#f8fafc"),
        "stroke": _hex(values.get("stroke", figure.design_tokens["line.data.color"]), "#334155"),
        "stroke_width": stroke_width_pt,
        "stroke_width_pt": stroke_width_pt,
        "stroke_width_mm": pt_to_mm(stroke_width_pt),
        "font_family": str(values.get("font_family", figure.design_tokens["font.family"])),
        "font_size": font_size_pt,
        "font_size_pt": font_size_pt,
        "font_size_mm": pt_to_mm(font_size_pt),
        "line_height_pt": float(values.get("line_height_pt", font_size_pt * 1.2)),
        "font_weight": int(values.get("font_weight", figure.design_tokens["font.weight"])),
        "opacity": max(0.0, min(1.0, float(values.get("opacity", figure.design_tokens["opacity"])))),
        "radius": radius_pt,
        "radius_pt": radius_pt,
        "radius_mm": pt_to_mm(radius_pt),
        "dash": str(values.get("dash", "")),
        "arrow": str(values.get("arrow", figure.design_tokens["arrow.kind"])),
        "arrow_size_pt": arrow_size_pt,
        "arrow_size_mm": pt_to_mm(arrow_size_pt),
    }


@dataclass(slots=True)
class CompiledPrimitive:
    kind: str
    object_id: str
    values: dict[str, Any]
    style: dict[str, Any]
    semantic: str


def _points(values: Iterable[Iterable[float]]) -> list[list[float]]:
    return [[round(float(x), 5), round(float(y), 5)] for x, y in values]


def _text_primitive(
    figure: FigureIR,
    item: FigureObject,
    *,
    text: str,
    role: str,
    style: dict[str, Any],
    semantic: str,
    box: tuple[float, float, float, float],
    horizontal_align: str = "start",
    vertical_align: str = "middle",
    real_shape: bool = False,
    full_text: str | None = None,
) -> CompiledPrimitive:
    layout = layout_text_box(
        text,
        x_mm=box[0],
        y_mm=box[1],
        width_mm=box[2],
        height_mm=box[3],
        font_size_pt=style["font_size_pt"],
        font_family=style["font_family"],
        font_weight=style["font_weight"],
        minimum_font_pt=float(figure.design_tokens["font.minimum.pt"]),
        line_height=style["line_height_pt"] / style["font_size_pt"],
        horizontal_align=horizontal_align,
        vertical_align=vertical_align,
    )
    if horizontal_align == "middle":
        measured_left = box[0] + (box[2] - layout.width_mm) / 2.0
    elif horizontal_align == "end":
        measured_left = box[0] + box[2] - layout.width_mm
    else:
        measured_left = box[0]
    values = {
        "x": layout.lines[0].x_mm if layout.lines else box[0],
        "y": layout.lines[0].baseline_y_mm if layout.lines else box[1],
        "text": text,
        "full_text": full_text or text,
        "lines": [
            {"text": line.text, "x": line.x_mm, "baseline_y": line.baseline_y_mm, "width_mm": line.width_mm}
            for line in layout.lines
        ],
        "text_box": {"x": box[0], "y": box[1], "width": box[2], "height": box[3]},
        "measured_bbox": {
            "x": measured_left,
            "y": layout.y_mm,
            "width": layout.width_mm,
            "height": layout.height_mm,
        },
        "anchor": horizontal_align,
        "vertical_anchor": vertical_align,
        "role": role,
        "real_shape": real_shape,
        "font_fallback": layout.font.to_dict(),
        "line_height_pt": layout.line_height_pt,
    }
    return CompiledPrimitive("text", item.id, values, style, semantic)


def _automatic_label_primitive(
    figure: FigureIR,
    item: FigureObject,
    *,
    label: str,
    role: str,
    style: dict[str, Any],
    semantic: str,
    box: tuple[float, float, float, float],
    horizontal_align: str,
    vertical_align: str,
) -> CompiledPrimitive:
    """Fit a generated label by abbreviation, never by glyph compression."""

    last_error: TextLayoutError | None = None
    tried: set[str] = set()
    for maximum_characters in range(max(2, min(28, len(label))), 1, -1):
        visible, _expanded = abbreviated_label(
            label,
            maximum_characters=maximum_characters,
        )
        if visible in tried:
            continue
        tried.add(visible)
        try:
            primitive = _text_primitive(
                figure,
                item,
                text=visible,
                full_text=label,
                role=role,
                style=style,
                semantic=semantic,
                box=box,
                horizontal_align=horizontal_align,
                vertical_align=vertical_align,
            )
        except TextLayoutError as exc:
            last_error = exc
            continue
        primitive.values["abbreviated"] = visible != label
        return primitive
    if last_error is not None:
        raise last_error
    raise TextLayoutError(f"Generated label {label!r} cannot fit its physical text lane")


def compile_object(
    figure: FigureIR,
    item: FigureObject,
    *,
    panel_geometry: dict[str, float] | None = None,
) -> list[CompiledPrimitive]:
    geometry, style = item.geometry, _style(figure, item)
    semantic = str(item.metadata.get("edge_semantics", item.provenance.kind))
    x, y = float(geometry.get("x", 0)), float(geometry.get("y", 0))
    width, height = float(geometry.get("width", 12)), float(geometry.get("height", 8))
    if item.kind == "tensor-glyph":
        author_shape = item.metadata.get("author_shape_override")
        evidence_backed_shape = (
            item.provenance.kind in {"graph_ir", "semantic_view"}
            and isinstance(item.provenance.evidence.get("tensor"), dict)
        )
        shape = list(author_shape if author_shape is not None else item.metadata.get("tensor_shape", []))
        shape_label = (
            item.metadata.get("author_shape_override_label")
            if author_shape is not None
            else item.metadata.get("shape_label", "")
        )
        mode = str(item.metadata.get("geometry_scale", "normalized"))
        manual = [width, height, float(geometry.get("depth", max(3.0, width * 0.2)))] if mode == "manual" else None
        visual = tensor_geometry(TensorGeometrySpec(
            shape,
            layout=str(item.metadata.get("layout", "auto")),
            scale_mode=mode,
            manual_size=manual,
            origin=(x, y),
            maximum_extent=max(8.0, min(42.0, max(width, height))),
            minimum_extent=max(3.5, min(width, height) * 0.25),
            label=str(shape_label),
            dtype=str(item.metadata.get("dtype", "unknown")),
        ))
        colors = {
            "front": _hex(item.style.overrides.get("front_fill", figure.design_tokens["color.tensor.front"]), "#dbeafe"),
            "top": _hex(item.style.overrides.get("top_fill", figure.design_tokens["color.tensor.top"]), "#bfdbfe"),
            "side": _hex(item.style.overrides.get("side_fill", figure.design_tokens["color.tensor.side"]), "#93c5fd"),
        }
        tensor_result = [
            CompiledPrimitive("polygon", item.id, {"points": visual.faces[face], "face": face, "fill": colors[face]}, style, semantic)
            for face in visual.draw_order
        ]
        lane_width = float(item.metadata.get("label_lane_width_mm", visual.bounds["width"] * 1.7))
        lane_center = visual.bounds["x"] + visual.bounds["width"] / 2.0
        shape_primitive = _text_primitive(
            figure,
            item,
            text=visual.shape_label,
            role="tensor-shape-label",
            style=style,
            semantic=semantic,
            box=(
                lane_center - lane_width / 2.0,
                visual.bounds["y"] + visual.bounds["height"] + pt_to_mm(style["font_size_pt"] * 0.3),
                lane_width,
                max(pt_to_mm(style["line_height_pt"]) * 3.0, 9.0),
            ),
            horizontal_align="middle",
            vertical_align="top",
            real_shape=author_shape is None and evidence_backed_shape,
        )
        if author_shape is not None:
            shape_origin = "author_override"
        elif evidence_backed_shape:
            shape_origin = "model_evidence"
        else:
            shape_origin = str(item.metadata.get("shape_label_origin") or (
                "author_annotation"
                if item.provenance.kind == "author_annotation"
                else "template_archetype"
            ))
        shape_primitive.values["shape_label_origin"] = shape_origin
        tensor_result.append(shape_primitive)
        return tensor_result
    if item.kind == "operator-glyph":
        symbol = operator_symbol(str(item.metadata.get("operator_family", item.metadata.get("op_type", "generic"))), x + width / 2, y + height / 2, min(width, height))
        operator_result: list[CompiledPrimitive] = []
        for value in symbol["primitives"]:
            primitive = dict(value)
            kind = str(primitive.pop("kind"))
            operator_result.append(CompiledPrimitive(kind, item.id, primitive, style, semantic))
        label = str(item.metadata.get("label", item.name))
        if label:
            lane_width = float(item.metadata.get("label_lane_width_mm", width * 1.2))
            label_box = (
                x + width / 2.0 - lane_width / 2.0,
                y + height + pt_to_mm(style["font_size_pt"] * 0.25),
                lane_width,
                pt_to_mm(style["line_height_pt"]) * 2.0,
            )
            if "short_label" in item.metadata:
                visible_label = str(item.metadata["short_label"])
                label_primitive = _text_primitive(
                    figure,
                    item,
                    text=visible_label,
                    full_text=label,
                    role="operator-label",
                    style=style,
                    semantic=semantic,
                    box=label_box,
                    horizontal_align="middle",
                    vertical_align="top",
                )
                label_primitive.values["abbreviated"] = visible_label != label
            else:
                label_primitive = _automatic_label_primitive(
                    figure,
                    item,
                    label=label,
                    role="operator-label",
                    style=style,
                    semantic=semantic,
                    box=label_box,
                    horizontal_align="middle",
                    vertical_align="top",
                )
            operator_result.append(label_primitive)
        return operator_result
    if item.kind in {"node-glyph", "image", "external-vector-group"}:
        node_result = [CompiledPrimitive("rect", item.id, {"x": x, "y": y, "width": width, "height": height, "rx": style["radius"]}, style, semantic)]
        if item.kind == "image" or item.kind == "external-vector-group":
            href = str(item.metadata.get("href", ""))
            if item.kind == "external-vector-group" and item.metadata.get("svg_source"):
                raw = str(item.metadata["svg_source"]).encode("utf-8")
                href = "data:image/svg+xml;base64," + base64.b64encode(raw).decode("ascii")
            if href:
                node_result.append(CompiledPrimitive("image", item.id, {"x": x, "y": y, "width": width, "height": height, "href": href}, style, semantic))
        else:
            padding = max(0.8, pt_to_mm(style["font_size_pt"] * 0.25))
            label_box = (
                x + padding,
                y + padding,
                max(0.1, width - 2 * padding),
                max(0.1, height - 2 * padding),
            )
            if "short_label" in item.metadata:
                visible_label = str(item.metadata["short_label"])
                label_primitive = _text_primitive(
                    figure,
                    item,
                    text=visible_label,
                    full_text=item.name,
                    role="node-label",
                    style=style,
                    semantic=semantic,
                    box=label_box,
                    horizontal_align="middle",
                    vertical_align="middle",
                )
                label_primitive.values["abbreviated"] = visible_label != item.name
            else:
                label_primitive = _automatic_label_primitive(
                    figure,
                    item,
                    label=item.name,
                    role="node-label",
                    style=style,
                    semantic=semantic,
                    box=label_box,
                    horizontal_align="middle",
                    vertical_align="middle",
                )
            node_result.append(label_primitive)
        return node_result
    if item.kind == "edge":
        route = item.manual_route or [[x, y], [float(geometry.get("x2", x + width)), float(geometry.get("y2", y + height))]]
        edge_style = dict(style)
        if semantic != "model-data":
            edge_style["stroke"] = _hex(item.style.overrides.get("stroke", figure.design_tokens["line.annotation.color"]), "#7c3aed")
            edge_style["dash"] = edge_style["dash"] or str(figure.design_tokens["line.annotation.dash"])
        return [CompiledPrimitive("polyline", item.id, {"points": _points(route), "arrow": edge_style["arrow"]}, edge_style, semantic)]
    if item.kind in {"region", "inset"}:
        region_style = dict(style)
        region_style["dash"] = region_style["dash"] or "4 3"
        return [CompiledPrimitive("rect", item.id, {"x": x, "y": y, "width": width, "height": height, "rx": style["radius"], "label": item.name}, region_style, semantic)]
    if item.kind in {"annotation", "title", "caption", "legend", "equation", "panel-label", "bracket"}:
        source_text = str(item.metadata.get("formula_source", item.metadata.get("text", item.name)))
        formula = normalize_formula(source_text) if item.kind == "equation" else None
        text = formula.text if formula is not None else source_text
        anchor = str(item.metadata.get("anchor", "start"))
        vertical_anchor = str(item.metadata.get("vertical_anchor", "bottom" if item.kind == "caption" else "middle"))
        if "width" in geometry:
            box_x, box_width = x, max(0.1, width)
        elif panel_geometry:
            box_x = x
            box_width = max(0.1, panel_geometry["x"] + panel_geometry["width"] - x - 3.0)
        else:
            box_x, box_width = x, float(item.metadata.get("max_width_mm", 80.0))
        box_height = float(geometry.get("height", item.metadata.get("max_height_mm", 20.0 if item.kind == "caption" else 12.0)))
        box_y = y - box_height if vertical_anchor == "bottom" and "height" not in geometry else y
        text_primitive = _text_primitive(
            figure,
            item,
            text=text,
            role=item.kind,
            style=style,
            semantic=semantic,
            box=(box_x, box_y, box_width, box_height),
            horizontal_align=anchor,
            vertical_align=vertical_anchor,
        )
        if formula is not None:
            # Formula output remains ordinary, editable text.  These values
            # are a complete audit record of the safe normalization decision;
            # no exporter is allowed to re-interpret or execute ``source``.
            text_primitive.values.update({
                "formula_source": formula.source,
                "formula_syntax_version": formula.syntax_version,
                "formula_valid": formula.valid,
                "formula_fallback": formula.fallback,
                "formula_error": formula.error,
                "formula_features": list(formula.features),
            })
        return [text_primitive]
    if item.kind in {"callout", "guide"}:
        route = item.manual_route or [[x, y], [float(geometry.get("x2", x + width)), float(geometry.get("y2", y + height))]]
        result = [CompiledPrimitive("polyline", item.id, {"points": _points(route), "arrow": style["arrow"]}, style, semantic)]
        if item.name:
            result.append(_text_primitive(
                figure,
                item,
                text=item.name,
                role=item.kind,
                style=style,
                semantic=semantic,
                box=(x, y - 1.5 - pt_to_mm(style["line_height_pt"]), float(item.metadata.get("max_width_mm", 50.0)), pt_to_mm(style["line_height_pt"])),
                horizontal_align="start",
                vertical_align="middle",
            ))
        return result
    raise ValidationError(f"Cannot compile Figure IR object kind {item.kind!r}")


def compile_figure(figure: FigureIR) -> list[tuple[str, str, str, CompiledPrimitive]]:
    figure.validate()
    compiled: list[tuple[str, str, str, CompiledPrimitive]] = []
    page_offset = 0.0
    for page in sorted(figure.pages, key=lambda item: (item.order, item.id)):
        for panel in sorted(page.panels, key=lambda item: (item.order, item.id)):
            for layer in sorted(panel.layers, key=lambda item: (item.order, item.id)):
                if not layer.visible:
                    continue
                for item in sorted(layer.objects, key=lambda value: (value.order, value.id)):
                    if not item.visible:
                        continue
                    for primitive in compile_object(figure, item, panel_geometry=panel.geometry):
                        if page_offset:
                            values = primitive.values
                            for key in ("y", "y1", "y2", "cy"):
                                if key in values:
                                    values[key] = float(values[key]) + page_offset
                            if "points" in values:
                                values["points"] = [[float(x), float(y) + page_offset] for x, y in values["points"]]
                            for line in values.get("lines", []):
                                line["baseline_y"] = float(line["baseline_y"]) + page_offset
                            for box_name in ("text_box", "measured_bbox"):
                                if box_name in values:
                                    values[box_name]["y"] = float(values[box_name]["y"]) + page_offset
                        compiled.append((page.id, panel.id, layer.id, primitive))
        page_offset += page.height_mm + 6.0
    return compiled


def _svg_style(style: dict[str, Any], *, fill: str | None = None) -> str:
    dash = ""
    if style.get("dash"):
        try:
            converted = " ".join(f"{pt_to_mm(float(value)):g}" for value in str(style["dash"]).replace(",", " ").split())
        except ValueError as exc:
            raise ValidationError(f"SVG dash pattern must contain point values: {style['dash']!r}") from exc
        dash = f' stroke-dasharray="{_esc(converted)}"'
    return (
        f' fill="{_esc(fill if fill is not None else style["fill"])}" '
        f'fill-opacity="{float(style["opacity"]):g}" stroke="{_esc(style["stroke"])}" '
        f'stroke-width="{float(style["stroke_width_mm"]):g}"{dash}'
    )


def render_figure_svg(
    figure: FigureIR,
    *,
    include_guides: bool = False,
    embed_metadata: bool = True,
) -> str:
    """Render a Figure IR document as physical-unit SVG.

    Persisted/exported SVG keeps the complete, digest-bound Figure IR metadata
    by default.  The editor preview may explicitly omit that duplicate payload
    because its JSON response already carries the authoritative Figure IR; this
    reduces interactive transfer and DOM parsing without changing export or
    SVG-import round trips.
    """

    figure.validate()
    width = max(page.width_mm for page in figure.pages)
    height = sum(page.height_mm for page in figure.pages) + max(0, len(figure.pages) - 1) * 6.0
    parts = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
        f'width="{width:g}mm" height="{height:g}mm" viewBox="0 0 {width:g} {height:g}" '
        f'data-nndv-document="figure-ir" data-nndv-version="{__version__}" data-figure-ir-version="{FIGURE_IR_VERSION}" '
        f'data-geometry-unit="mm" data-style-unit="pt" data-css-dpi="96">',
        f'<title>{_esc(figure.name)}</title>',
    ]
    if embed_metadata:
        metadata = {
            "schema_version": FIGURE_SVG_METADATA_VERSION,
            "nn_davinci_version": __version__,
            "figure_ir_version": FIGURE_IR_VERSION,
            "figure_digest": figure.digest(),
            "figure_ir": figure.to_dict(),
        }
        parts.append(
            f'<metadata id="nndv-figure-ir">{_esc(json.dumps(metadata, ensure_ascii=False, sort_keys=True, separators=(",", ":")))}</metadata>'
        )
    parts.extend(
        [
            (
                '<defs><marker id="nndv-arrow" markerUnits="userSpaceOnUse" '
                f'markerWidth="{pt_to_mm(float(figure.design_tokens["arrow.size.pt"]) * 1.5):g}" '
                f'markerHeight="{pt_to_mm(float(figure.design_tokens["arrow.size.pt"])):g}" '
                'refX="6" refY="2.5" orient="auto" viewBox="0 0 6 5" preserveAspectRatio="xMidYMid meet">'
                '<path d="M0,0 L6,2.5 L0,5 Z" fill="context-stroke"/></marker></defs>'
            ),
            f'<rect class="figure-background" width="100%" height="100%" fill="{_esc(figure.design_tokens["color.background"])}"/>',
        ]
    )
    page_offset = 0.0
    primitive_map: dict[tuple[str, str, str], list[CompiledPrimitive]] = {}
    for page_id, panel_id, layer_id, primitive in compile_figure(figure):
        primitive_map.setdefault((page_id, panel_id, layer_id), []).append(primitive)
    for page in sorted(figure.pages, key=lambda item: (item.order, item.id)):
        parts.append(f'<g class="figure-page" data-page-id="{_esc(page.id)}" transform="translate(0,{page_offset:g})">')
        parts.append(f'<rect x="0" y="0" width="{page.width_mm:g}" height="{page.height_mm:g}" fill="none" stroke="none"/>')
        if include_guides:
            margin = page.margin_mm
            parts.append(f'<rect class="page-margin-guide" x="{margin:g}" y="{margin:g}" width="{page.width_mm-2*margin:g}" height="{page.height_mm-2*margin:g}" fill="none" stroke="{_esc(figure.design_tokens["guide.margin.color"])}" stroke-dasharray="2 2"/>')
            column_width = (page.width_mm - 2 * margin - (page.columns - 1) * page.column_gap_mm) / page.columns
            for index in range(1, page.columns):
                gx = margin + index * column_width + (index - 0.5) * page.column_gap_mm
                parts.append(f'<line class="column-guide" x1="{gx:g}" y1="{margin:g}" x2="{gx:g}" y2="{page.height_mm-margin:g}" stroke="{_esc(figure.design_tokens["guide.column.color"])}"/>')
            baseline_mm = page.baseline_grid_pt * 25.4 / 72.0
            baseline_y = margin + baseline_mm
            while baseline_y < page.height_mm - margin:
                parts.append(
                    f'<line class="baseline-guide" x1="{margin:g}" y1="{baseline_y:g}" '
                    f'x2="{page.width_mm-margin:g}" y2="{baseline_y:g}" '
                    f'stroke="{_esc(figure.design_tokens["guide.baseline.color"])}" stroke-width="{pt_to_mm(0.2):g}"/>'
                )
                baseline_y += baseline_mm
        for panel in sorted(page.panels, key=lambda item: (item.order, item.id)):
            box = panel.geometry
            parts.append(
                f'<g class="figure-panel panel-mode-{_esc(panel.mode)}" data-panel-id="{_esc(panel.id)}" '
                f'data-panel-label="{_esc(panel.label)}" data-panel-mode="{_esc(panel.mode)}" data-locked="{str(panel.locked).lower()}">'
            )
            parts.append(
                f'<rect class="panel-boundary" data-nndv-role="panel" data-panel-id="{_esc(panel.id)}" '
                f'x="{box["x"]:g}" y="{box["y"]:g}" width="{box["width"]:g}" height="{box["height"]:g}" '
                f'fill="none" stroke="#cbd5e1" stroke-width="{pt_to_mm(0.35):g}"/>'
            )
            panel_font_pt = max(7.0, float(figure.design_tokens["font.size.pt"]))
            panel_font_mm = pt_to_mm(panel_font_pt)
            panel_label_y = box["y"] + 1.2 + panel_font_mm * 0.8
            parts.append(
                f'<text class="panel-label" data-nndv-role="required-text" data-text-role="panel-label" '
                f'data-panel-id="{_esc(panel.id)}" data-font-size-pt="{panel_font_pt:g}" '
                f'x="{box["x"]+1.5:g}" y="{panel_label_y:g}" font-family="{_esc(figure.design_tokens["font.family"])}" '
                f'font-size="{panel_font_mm:g}" font-weight="700">{_esc(panel.label)}</text>'
            )
            for layer in sorted(panel.layers, key=lambda item: (item.order, item.id)):
                if not layer.visible:
                    continue
                parts.append(
                    f'<g class="figure-layer layer-{_esc(layer.role)}" data-layer-id="{_esc(layer.id)}" '
                    f'data-layer-role="{_esc(layer.role)}" data-locked="{str(layer.locked).lower()}">'
                )
                for primitive in primitive_map.get((page.id, panel.id, layer.id), []):
                    # compile_figure already applies the page offset; remove it
                    # inside the translated page group.
                    values = dict(primitive.values)
                    if page_offset:
                        for key in ("y", "y1", "y2", "cy"):
                            if key in values:
                                values[key] = float(values[key]) - page_offset
                        if "points" in values:
                            values["points"] = [[float(x), float(y) - page_offset] for x, y in values["points"]]
                        for line in values.get("lines", []):
                            line["baseline_y"] = float(line["baseline_y"]) - page_offset
                        for box_name in ("text_box", "measured_bbox"):
                            if box_name in values:
                                values[box_name]["y"] = float(values[box_name]["y"]) - page_offset
                    object_id = _esc(primitive.object_id)
                    role = "required-text" if primitive.kind == "text" else ("edge" if primitive.kind in {"polyline", "line", "path"} else "figure-object")
                    common = (
                        f'data-figure-object-id="{object_id}" data-object-semantic="{_esc(primitive.semantic)}" '
                        f'data-nndv-role="{role}" data-primitive-kind="{_esc(primitive.kind)}" '
                        f'data-page-id="{_esc(page.id)}" data-panel-id="{_esc(panel.id)}" data-layer-id="{_esc(layer.id)}"'
                    )
                    if primitive.kind == "rect":
                        parts.append(f'<rect {common} x="{values["x"]:g}" y="{values["y"]:g}" width="{values["width"]:g}" height="{values["height"]:g}" rx="{float(primitive.style["radius_mm"]):g}"{_svg_style(primitive.style)}/>')
                    elif primitive.kind == "polygon":
                        polygon_points_text = " ".join(f"{x:g},{y:g}" for x, y in values["points"])
                        parts.append(f'<polygon {common} data-tensor-face="{_esc(values.get("face", ""))}" points="{polygon_points_text}"{_svg_style(primitive.style, fill=values.get("fill"))}/>')
                    elif primitive.kind in {"polyline", "line"}:
                        if primitive.kind == "line":
                            line_points = [[values["x1"], values["y1"]], [values["x2"], values["y2"]]]
                        else:
                            line_points = values["points"]
                        point_text = " ".join(f"{x:g},{y:g}" for x, y in line_points)
                        marker = ' marker-end="url(#nndv-arrow)"' if values.get("arrow") in {"end", "both", "true"} else ""
                        marker += ' marker-start="url(#nndv-arrow)"' if values.get("arrow") in {"start", "both"} else ""
                        parts.append(f'<polyline {common} points="{point_text}"{_svg_style(primitive.style, fill="none")}{marker}/>')
                    elif primitive.kind == "circle":
                        parts.append(f'<circle {common} cx="{values["cx"]:g}" cy="{values["cy"]:g}" r="{values["r"]:g}"{_svg_style(primitive.style)}/>')
                    elif primitive.kind == "path":
                        parts.append(f'<path {common} d="{_esc(values["d"])}"{_svg_style(primitive.style, fill="none")}/>')
                    elif primitive.kind == "text":
                        real_shape_attr = ' data-real-shape="true"' if values.get("real_shape") else ""
                        shape_origin_attr = (
                            f' data-shape-label-origin="{_esc(values["shape_label_origin"])}"'
                            if values.get("shape_label_origin") else ""
                        )
                        abbreviated_attr = (
                            f' data-abbreviated="{str(bool(values["abbreviated"])).lower()}"'
                            if "abbreviated" in values else ""
                        )
                        formula_attr = ""
                        if "formula_source" in values:
                            formula_attr = (
                                f' data-formula-source="{_esc(values["formula_source"])}"'
                                f' data-formula-syntax-version="{_esc(values["formula_syntax_version"])}"'
                                f' data-formula-valid="{str(bool(values["formula_valid"])).lower()}"'
                                f' data-formula-fallback="{str(bool(values["formula_fallback"])).lower()}"'
                                f' data-formula-error="{_esc(values["formula_error"])}"'
                                f' data-formula-features="{_esc(json.dumps(values["formula_features"], ensure_ascii=False, separators=(",", ":")))}"'
                                ' data-formula-rendering="editable-text"'
                            )
                        fallback = values.get("font_fallback", {})
                        requested = ", ".join(fallback.get("requested", ()))
                        candidates = ", ".join(fallback.get("fallback_candidates", ()))
                        lines = list(values.get("lines") or [{"text": values["text"], "x": values["x"], "baseline_y": values["y"]}])
                        tspans = "".join(
                            f'<tspan x="{float(line["x"]):g}" y="{float(line["baseline_y"]):g}">{_esc(line["text"])}</tspan>'
                            for line in lines
                        )
                        parts.append(
                            f'<text {common} class="figure-text {_esc(values.get("role", "figure-text"))}" '
                            f'data-text-role="{_esc(values.get("role", "figure-text"))}" data-font-size-pt="{primitive.style["font_size_pt"]:g}" '
                            f'data-line-height-pt="{float(values.get("line_height_pt", primitive.style["line_height_pt"])):g}" '
                            f'data-font-requested="{_esc(requested)}" data-font-fallback-candidates="{_esc(candidates)}" '
                            f'data-full-text="{_esc(values.get("full_text", values["text"]))}" '
                            f'x="{float(values["x"]):g}" y="{float(values["y"]):g}" '
                            f'text-anchor="{_esc(values.get("anchor", "start"))}" xml:space="preserve" '
                            f'font-family="{_esc(fallback.get("css_stack", primitive.style["font_family"]))}" '
                            f'font-size="{primitive.style["font_size_mm"]:g}" font-weight="{primitive.style["font_weight"]}" '
                            f'fill="{_esc(figure.design_tokens["color.foreground"])}" fill-opacity="{float(primitive.style["opacity"]):g}"'
                            f'{real_shape_attr}{shape_origin_attr}{abbreviated_attr}{formula_attr}>{tspans}</text>'
                        )
                    elif primitive.kind == "image":
                        parts.append(f'<image {common} x="{values["x"]:g}" y="{values["y"]:g}" width="{values["width"]:g}" height="{values["height"]:g}" href="{_esc(values["href"])}" preserveAspectRatio="xMidYMid meet"/>')
                parts.append("</g>")
            parts.append("</g>")
        parts.append("</g>")
        page_offset += page.height_mm + 6.0
    parts.append("</svg>")
    return "\n".join(parts) + "\n"


def render_figure_svg_pages(figure: FigureIR, *, include_guides: bool = False) -> list[tuple[str, str]]:
    """Render page-sized SVG documents in canonical FigurePage order.

    ``render_figure_svg`` remains the backwards-compatible stacked document
    renderer used by the editor.  Exporters use this page-first projection so
    physical page size and orientation are never inferred from a contact sheet.
    Each SVG embeds the projected page plus a link to the complete source
    figure digest in ``metadata.page_export``.
    """

    figure.validate()
    return [
        (page.id, render_figure_svg(_page_figure(figure, index), include_guides=include_guides))
        for index, page in enumerate(_ordered_pages(figure))
    ]


def import_figure_svg(source: str | bytes, *, name: str = "Imported SVG") -> tuple[FigureIR, bool]:
    raw = source.decode("utf-8") if isinstance(source, bytes) else str(source)
    if len(raw.encode("utf-8")) > 16 * 1024 * 1024:
        raise ValidationError("SVG import exceeds the 16 MiB limit")
    lowered = raw.lower()
    if "<!doctype" in lowered or "<!entity" in lowered:
        raise ValidationError("SVG DTD and entity declarations are not accepted")
    try:
        root = ET.fromstring(raw)
    except ET.ParseError as exc:
        raise ValidationError("SVG XML is malformed") from exc
    if root.tag.rsplit("}", 1)[-1] != "svg":
        raise ValidationError("Imported XML root is not SVG")
    metadata_node = next((item for item in root.iter() if item.tag.rsplit("}", 1)[-1] == "metadata" and item.attrib.get("id") == "nndv-figure-ir"), None)
    if metadata_node is not None and metadata_node.text:
        try:
            metadata = json.loads(metadata_node.text)
        except json.JSONDecodeError as exc:
            raise ValidationError("NN_DaVinci SVG metadata is malformed") from exc
        if metadata.get("schema_version") != FIGURE_SVG_METADATA_VERSION:
            raise ValidationError("NN_DaVinci SVG metadata version is unsupported")
        figure = FigureIR.from_dict(metadata["figure_ir"])
        if metadata.get("figure_digest") != figure.digest():
            raise ValidationError("NN_DaVinci SVG Figure IR digest does not match its metadata")
        return figure, True
    figure = new_figure(name)
    panel = next(figure.iter_panels())
    layer = next(item for item in panel.layers if item.role == "author-annotation")
    layer.objects.append(FigureObject.create(
        "external-vector-group", name,
        {"x": panel.geometry["x"] + 4, "y": panel.geometry["y"] + 7, "width": panel.geometry["width"] - 8, "height": panel.geometry["height"] - 11},
        FigureProvenance("external_vector", reason="Third-party SVG imported as an external vector group; no model semantics were recovered."),
        identity=f"external-svg:{len(raw)}:{raw[:120]}",
        metadata={"svg_source": raw, "semantic_recovery": False, "external_vector": True},
    ))
    figure.metadata["import"] = {"kind": "third-party-svg", "model_semantics_recovered": False}
    return figure.validate(), False


_TEX_EDITABLE_UNICODE = {
    # Greek produced by the safe formula normalizer.
    "α": r"\ensuremath{\alpha}", "β": r"\ensuremath{\beta}", "γ": r"\ensuremath{\gamma}",
    "δ": r"\ensuremath{\delta}", "ε": r"\ensuremath{\epsilon}", "ϵ": r"\ensuremath{\varepsilon}",
    "ζ": r"\ensuremath{\zeta}", "η": r"\ensuremath{\eta}", "θ": r"\ensuremath{\theta}",
    "ϑ": r"\ensuremath{\vartheta}", "ι": r"\ensuremath{\iota}", "κ": r"\ensuremath{\kappa}",
    "λ": r"\ensuremath{\lambda}", "μ": r"\ensuremath{\mu}", "ν": r"\ensuremath{\nu}",
    "ξ": r"\ensuremath{\xi}", "π": r"\ensuremath{\pi}", "ϖ": r"\ensuremath{\varpi}",
    "ρ": r"\ensuremath{\rho}", "σ": r"\ensuremath{\sigma}", "τ": r"\ensuremath{\tau}",
    "υ": r"\ensuremath{\upsilon}", "φ": r"\ensuremath{\phi}", "ϕ": r"\ensuremath{\varphi}",
    "χ": r"\ensuremath{\chi}", "ψ": r"\ensuremath{\psi}", "ω": r"\ensuremath{\omega}",
    "Γ": r"\ensuremath{\Gamma}", "Δ": r"\ensuremath{\Delta}", "Θ": r"\ensuremath{\Theta}",
    "Λ": r"\ensuremath{\Lambda}", "Ξ": r"\ensuremath{\Xi}", "Π": r"\ensuremath{\Pi}",
    "Σ": r"\ensuremath{\Sigma}", "Υ": r"\ensuremath{\Upsilon}", "Φ": r"\ensuremath{\Phi}",
    "Ψ": r"\ensuremath{\Psi}", "Ω": r"\ensuremath{\Omega}",
    # Operators and scientific symbols produced by normalization.
    "×": r"\ensuremath{\times}", "·": r"\ensuremath{\cdot}", "±": r"\ensuremath{\pm}",
    "∓": r"\ensuremath{\mp}", "≤": r"\ensuremath{\leq}", "≥": r"\ensuremath{\geq}",
    "≠": r"\ensuremath{\neq}", "≈": r"\ensuremath{\approx}", "∼": r"\ensuremath{\sim}",
    "→": r"\ensuremath{\rightarrow}", "←": r"\ensuremath{\leftarrow}", "↔": r"\ensuremath{\leftrightarrow}",
    "∞": r"\ensuremath{\infty}", "∑": r"\ensuremath{\sum}", "∏": r"\ensuremath{\prod}",
    "∫": r"\ensuremath{\int}", "∂": r"\ensuremath{\partial}", "∇": r"\ensuremath{\nabla}",
    "∈": r"\ensuremath{\in}", "∉": r"\ensuremath{\notin}", "⊂": r"\ensuremath{\subset}",
    "⊆": r"\ensuremath{\subseteq}", "∪": r"\ensuremath{\cup}", "∩": r"\ensuremath{\cap}",
    "∀": r"\ensuremath{\forall}", "∃": r"\ensuremath{\exists}", "¬": r"\ensuremath{\neg}",
    "∧": r"\ensuremath{\land}", "∨": r"\ensuremath{\lor}", "ℓ": r"\ensuremath{\ell}",
    "ℜ": r"\ensuremath{\Re}", "ℑ": r"\ensuremath{\Im}", "ℝ": r"\ensuremath{\mathbb{R}}",
    "ℕ": r"\ensuremath{\mathbb{N}}", "√": r"\ensuremath{\surd}",
    # Unicode scripts remain independent, editable text runs in TikZ.
    "⁰": r"\textsuperscript{0}", "¹": r"\textsuperscript{1}", "²": r"\textsuperscript{2}",
    "³": r"\textsuperscript{3}", "⁴": r"\textsuperscript{4}", "⁵": r"\textsuperscript{5}",
    "⁶": r"\textsuperscript{6}", "⁷": r"\textsuperscript{7}", "⁸": r"\textsuperscript{8}",
    "⁹": r"\textsuperscript{9}", "⁺": r"\textsuperscript{+}", "⁻": r"\textsuperscript{-}",
    "⁼": r"\textsuperscript{=}", "⁽": r"\textsuperscript{(}", "⁾": r"\textsuperscript{)}",
    "ⁱ": r"\textsuperscript{i}", "ʲ": r"\textsuperscript{j}", "ᵏ": r"\textsuperscript{k}",
    "ᵀ": r"\textsuperscript{T}",
    "₀": r"\textsubscript{0}", "₁": r"\textsubscript{1}", "₂": r"\textsubscript{2}",
    "₃": r"\textsubscript{3}", "₄": r"\textsubscript{4}", "₅": r"\textsubscript{5}",
    "₆": r"\textsubscript{6}", "₇": r"\textsubscript{7}", "₈": r"\textsubscript{8}",
    "₉": r"\textsubscript{9}", "₊": r"\textsubscript{+}", "₋": r"\textsubscript{-}",
    "₌": r"\textsubscript{=}", "₍": r"\textsubscript{(}", "₎": r"\textsubscript{)}",
    "ₐ": r"\textsubscript{a}", "ₑ": r"\textsubscript{e}", "ₒ": r"\textsubscript{o}",
    "ₓ": r"\textsubscript{x}", "ₕ": r"\textsubscript{h}", "ₖ": r"\textsubscript{k}",
    "ₗ": r"\textsubscript{l}", "ₘ": r"\textsubscript{m}", "ₙ": r"\textsubscript{n}",
    "ₚ": r"\textsubscript{p}", "ₛ": r"\textsubscript{s}", "ₜ": r"\textsubscript{t}",
    "ᵢ": r"\textsubscript{i}", "ⱼ": r"\textsubscript{j}", "ᵣ": r"\textsubscript{r}",
}


def _tex(value: Any) -> str:
    """Serialize editable text without ever interpreting it as TeX source."""

    special = {
        "\\": r"\textbackslash{}", "{": r"\{", "}": r"\}", "$": r"\$",
        "_": r"\_", "%": r"\%", "&": r"\&", "#": r"\#",
        "^": r"\textasciicircum{}", "~": r"\textasciitilde{}",
    }
    return "".join(_TEX_EDITABLE_UNICODE.get(character, special.get(character, character)) for character in str(value))


def render_figure_tikz(figure: FigureIR) -> str:
    figure.validate()
    width = max(page.width_mm for page in figure.pages)
    height = sum(page.height_mm for page in figure.pages) + max(0, len(figure.pages) - 1) * 6.0
    compiled = compile_figure(figure)
    colours = {_hex(figure.design_tokens["color.foreground"]), _hex("#cbd5e1")}
    for _, _, _, primitive in compiled:
        colours.add(_hex(primitive.style.get("fill")))
        colours.add(_hex(primitive.style.get("stroke")))
        if primitive.values.get("fill"):
            colours.add(_hex(primitive.values["fill"]))
    colour_names = {colour: f"nndvcolor{index}" for index, colour in enumerate(sorted(colours))}
    lines = [
        r"\documentclass[tikz,border=0pt]{standalone}",
        r"\usepackage[T1]{fontenc}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage{lmodern}",
        r"\usepackage{amsmath,amssymb}",
        r"\usepackage{tikz}",
        *(
            rf"\definecolor{{{name}}}{{HTML}}{{{colour[1:].upper()}}}"
            for colour, name in colour_names.items()
        ),
        r"\begin{document}",
        rf"\begin{{tikzpicture}}[x=1mm,y=-1mm,font=\sffamily\fontsize{{{float(figure.design_tokens['font.size.pt']):g}}}{{{float(figure.design_tokens['font.size.pt'])*1.2:g}}}\selectfont]",
        rf"\path[use as bounding box] (0,0) rectangle ({width:g},{height:g});",
    ]
    page_offset = 0.0
    panel_colour = colour_names[_hex("#cbd5e1")]
    text_colour = colour_names[_hex(figure.design_tokens["color.foreground"])]
    panel_font_pt = max(7.0, float(figure.design_tokens["font.size.pt"]))
    for page in sorted(figure.pages, key=lambda item: (item.order, item.id)):
        for panel in sorted(page.panels, key=lambda item: (item.order, item.id)):
            box = panel.geometry
            top = float(box["y"]) + page_offset
            lines.append(f"% NNDV panel {panel.id} {panel.label}")
            lines.append(
                rf"\draw[draw={panel_colour},line width=0.35pt] ({box['x']:g},{top:g}) "
                rf"rectangle ({box['x'] + box['width']:g},{top + box['height']:g});"
            )
            label_y = top + 1.2 + pt_to_mm(panel_font_pt) * 0.8
            lines.append(
                rf"\node[anchor=base west,text={text_colour},inner sep=0pt,"
                rf"font=\sffamily\bfseries\fontsize{{{panel_font_pt:g}}}{{{panel_font_pt * 1.2:g}}}\selectfont] "
                rf"at ({box['x'] + 1.5:g},{label_y:g}) {{{_tex(panel.label)}}};"
            )
        page_offset += page.height_mm + 6.0
    for _, panel_id, layer_id, primitive in compiled:
        values, style = primitive.values, primitive.style
        color = colour_names[_hex(values.get("fill", style["fill"]))]
        stroke = colour_names[_hex(style["stroke"])]
        dash = ",dashed" if style.get("dash") else ""
        face_marker = f" tensor-face={values['face']}" if values.get("face") else ""
        comment = f"% NNDV {panel_id} {layer_id} {primitive.object_id} {primitive.semantic} {primitive.kind}{face_marker}"
        lines.append(comment)
        if primitive.kind == "rect":
            lines.append(rf"\draw[draw={stroke},fill={color},line width={style['stroke_width']:g}pt{dash},rounded corners={float(values.get('rx',0)):g}pt] ({values['x']:g},{values['y']:g}) rectangle ({values['x']+values['width']:g},{values['y']+values['height']:g});")
        elif primitive.kind == "polygon":
            point_text = " -- ".join(f"({x:g},{y:g})" for x, y in values["points"])
            polygon_fill = colour_names[_hex(values.get("fill", style["fill"]))]
            lines.append(rf"\draw[draw={stroke},fill={polygon_fill},line width={style['stroke_width']:g}pt{dash}] {point_text} -- cycle;")
        elif primitive.kind in {"polyline", "line"}:
            points = values.get("points") or [[values["x1"], values["y1"]], [values["x2"], values["y2"]]]
            point_text = " -- ".join(f"({x:g},{y:g})" for x, y in points)
            arrow = "->" if values.get("arrow") in {"end", "both", "true"} else "-"
            if values.get("arrow") == "both":
                arrow = "<->"
            lines.append(rf"\draw[{arrow},draw={stroke},line width={style['stroke_width']:g}pt{dash}] {point_text};")
        elif primitive.kind == "circle":
            lines.append(rf"\draw[draw={stroke},fill={color},line width={style['stroke_width']:g}pt] ({values['cx']:g},{values['cy']:g}) circle ({values['r']:g}mm);")
        elif primitive.kind == "text":
            anchor = {"middle": "base", "end": "base east"}.get(str(values.get("anchor", "")), "base west")
            weight = r"\bfseries" if int(style["font_weight"]) >= 600 else ""
            font = (
                rf"\sffamily{weight}\fontsize{{{style['font_size_pt']:g}}}"
                rf"{{{float(values.get('line_height_pt', style['line_height_pt'])):g}}}\selectfont"
            )
            text_lines = values.get("lines") or [{"text": values["text"], "x": values["x"], "baseline_y": values["y"]}]
            for text_line in text_lines:
                lines.append(
                    rf"\node[anchor={anchor},text={text_colour},inner sep=0pt,font={font}] "
                    rf"at ({float(text_line['x']):g},{float(text_line['baseline_y']):g}) "
                    rf"{{{_tex(text_line['text'])}}};"
                )
    lines.extend([r"\end{tikzpicture}", r"\end{document}"])
    return "\n".join(lines) + "\n"


def render_figure_tikz_pages(figure: FigureIR) -> list[tuple[str, str]]:
    """Render one standalone, page-sized TikZ source per FigurePage."""

    figure.validate()
    return [
        (page.id, render_figure_tikz(_page_figure(figure, index)))
        for index, page in enumerate(_ordered_pages(figure))
    ]


def _load_cairosvg() -> Any:
    try:
        import cairosvg
    except ImportError as exc:
        raise OptionalDependencyError(
            "Figure PDF, PNG, and EPS export requires CairoSVG; install the project's locked export dependencies"
        ) from exc
    return cairosvg


def _run_converter(arguments: list[str], *, purpose: str) -> None:
    try:
        result = subprocess.run(arguments, check=False, capture_output=True, text=True)
    except OSError as exc:
        raise OptionalDependencyError(f"{purpose} could not start {arguments[0]!r}: {exc}") from exc
    if result.returncode:
        detail = (result.stderr or result.stdout or "no diagnostic output").strip()
        raise ExportError(f"{purpose} failed with exit code {result.returncode}: {detail}")


def export_figure_pdf(figure: FigureIR, destination: str | Path) -> Path:
    """Write a truly paginated PDF while retaining each page's physical size."""

    figure.validate()
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    page_svgs = render_figure_svg_pages(figure)
    cairosvg = _load_cairosvg()
    if len(page_svgs) == 1:
        try:
            cairosvg.svg2pdf(bytestring=page_svgs[0][1].encode("utf-8"), write_to=str(destination))
        except Exception as exc:
            raise ExportError(f"Figure PDF page rendering failed: {exc}") from exc
    else:
        pdfunite = shutil.which("pdfunite")
        if pdfunite is None:
            raise OptionalDependencyError(
                "Multi-page Figure PDF export requires the offline Poppler 'pdfunite' executable; "
                "install poppler-utils or request per-page SVG export"
            )
        with tempfile.TemporaryDirectory(prefix="nndv-figure-pdf-") as directory:
            page_pdfs: list[Path] = []
            for index, (_page_id, svg) in enumerate(page_svgs):
                page_pdf = Path(directory) / f"page-{index + 1:03d}.pdf"
                try:
                    cairosvg.svg2pdf(bytestring=svg.encode("utf-8"), write_to=str(page_pdf))
                except Exception as exc:
                    raise ExportError(f"Figure PDF page {index + 1} rendering failed: {exc}") from exc
                page_pdfs.append(page_pdf)
            _run_converter(
                [pdfunite, *(str(path) for path in page_pdfs), str(destination)],
                purpose="Multi-page Figure PDF assembly",
            )
    if not destination.is_file() or not destination.read_bytes().startswith(b"%PDF"):
        raise ExportError("Figure PDF exporter did not produce a valid PDF signature")
    return destination


def export_figure_png_pages(figure: FigureIR, stem: str | Path) -> list[Path]:
    """Write one 300-DPI, metadata-bearing RGBA PNG for each page."""

    figure.validate()
    try:
        from PIL import Image, PngImagePlugin
    except ImportError as exc:
        raise OptionalDependencyError(
            "Figure PNG metadata export requires Pillow; install the project's locked export dependencies"
        ) from exc
    cairosvg = _load_cairosvg()
    stem = Path(stem).with_suffix("")
    stem.parent.mkdir(parents=True, exist_ok=True)
    pages = _ordered_pages(figure)
    page_svgs = render_figure_svg_pages(figure)
    policy = figure_export_policy(figure)
    outputs: list[Path] = []
    for index, ((page_id, svg), page) in enumerate(zip(page_svgs, pages, strict=True)):
        width_px = max(1, round(page.width_mm / 25.4 * PNG_EXPORT_DPI))
        height_px = max(1, round(page.height_mm / 25.4 * PNG_EXPORT_DPI))
        try:
            rendered = cairosvg.svg2png(
                bytestring=svg.encode("utf-8"), output_width=width_px, output_height=height_px,
            )
            with Image.open(io.BytesIO(rendered)) as opened:
                image = opened.convert("RGBA")
        except Exception as exc:
            raise ExportError(f"Figure PNG page {index + 1} rendering failed: {exc}") from exc
        metadata = {
            "schema_version": "nndv-png-export-1",
            "source_figure_digest": figure.digest(),
            "page_id": page_id,
            "page_index": index,
            "page_number": index + 1,
            "page_count": len(pages),
            "width_mm": page.width_mm,
            "height_mm": page.height_mm,
            "pixel_width": width_px,
            "pixel_height": height_px,
            **policy["formats"]["png"],
        }
        png_info = PngImagePlugin.PngInfo()
        png_info.add_text("Software", "NN_DaVinci Figure Export")
        png_info.add_text("nndv.figure_export", json.dumps(metadata, ensure_ascii=False, sort_keys=True, separators=(",", ":")))
        destination = _page_destination(stem, ".png", index, len(pages))
        image.save(destination, format="PNG", dpi=(PNG_EXPORT_DPI, PNG_EXPORT_DPI), pnginfo=png_info)
        outputs.append(destination)
    return outputs


def export_figure_eps_pages(figure: FigureIR, stem: str | Path) -> list[Path]:
    """Write vector EPSF pages through Poppler, rejecting raster fallbacks."""

    figure.validate()
    pdftops = shutil.which("pdftops")
    if pdftops is None:
        raise OptionalDependencyError(
            "Vector EPS export requires the offline Poppler 'pdftops' executable; "
            "NN_DaVinci will not silently rasterize EPS output"
        )
    cairosvg = _load_cairosvg()
    stem = Path(stem).with_suffix("")
    stem.parent.mkdir(parents=True, exist_ok=True)
    pages = _ordered_pages(figure)
    page_svgs = render_figure_svg_pages(figure)
    outputs: list[Path] = []
    with tempfile.TemporaryDirectory(prefix="nndv-figure-eps-") as directory:
        for index, ((page_id, svg), page) in enumerate(zip(page_svgs, pages, strict=True)):
            page_pdf = Path(directory) / f"page-{index + 1:03d}.pdf"
            try:
                cairosvg.svg2pdf(bytestring=svg.encode("utf-8"), write_to=str(page_pdf))
            except Exception as exc:
                raise ExportError(f"Figure EPS page {index + 1} vector staging failed: {exc}") from exc
            destination = _page_destination(stem, ".eps", index, len(pages))
            _run_converter(
                [pdftops, "-eps", "-level3", "-f", "1", "-l", "1", str(page_pdf), str(destination)],
                purpose=f"Figure EPS page {index + 1} vector conversion",
            )
            payload = destination.read_bytes()
            if not payload.startswith(b"%!PS-Adobe-3.0 EPSF-3.0"):
                raise ExportError(f"Figure EPS page {index + 1} is not an EPSF Level 3 vector file")
            if b"%%BeginResource: font" not in payload:
                raise ExportError(
                    f"Figure EPS page {index + 1} has no embedded font resources; refusing an unverified export"
                )
            first_line, separator, remainder = payload.partition(b"\n")
            dsc = (
                f"%%NNDVSourceFigureDigest: {figure.digest()}\n"
                f"%%NNDVPage: {index + 1} {len(pages)}\n"
                f"%%NNDVPageId: {page_id}\n"
                f"%%NNDVPageSizeMM: {page.width_mm:g} {page.height_mm:g}\n"
                "%%NNDVVectorPolicy: SVG-to-PDF-to-Poppler-EPS; no page rasterization fallback\n"
                "%%NNDVFontPolicy: embedded subset resources required\n"
            ).encode("ascii")
            destination.write_bytes(first_line + separator + dsc + remainder)
            outputs.append(destination)
    return outputs


def _safe_script_json(value: Any) -> str:
    return (
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        .replace("&", "\\u0026")
        .replace("<", "\\u003c")
        .replace(">", "\\u003e")
    )


def _validate_offline_html_assets(figure: FigureIR) -> None:
    for item in figure.iter_objects(visible_only=True):
        if item.kind != "image":
            continue
        href = str(item.metadata.get("href", ""))
        if href and not href.startswith("data:"):
            raise ExportError(
                f"Offline HTML export cannot fetch external image {href!r} for object {item.id!r}; "
                "embed it as a data URI first"
            )


def render_figure_html(figure: FigureIR) -> str:
    """Return a dependency-free, source-preserving, editable page viewer."""

    figure.validate()
    _validate_offline_html_assets(figure)
    pages = _ordered_pages(figure)
    page_svgs = render_figure_svg_pages(figure)
    source_json = _safe_script_json(figure.to_dict())
    policy_json = _safe_script_json(figure_export_policy(figure))
    page_options = "".join(
        f'<option value="{index}">{index + 1}: {_esc(page.title)}</option>'
        for index, page in enumerate(pages)
    )
    page_sections: list[str] = []
    for index, ((page_id, svg), page) in enumerate(zip(page_svgs, pages, strict=True)):
        inline_svg = svg.split("?>", 1)[-1].lstrip()
        # HTML IDs are document-global, unlike IDs in separate SVG files.
        # Namespace the two fixed Figure SVG IDs so multi-page viewers remain
        # valid DOM documents and markers never resolve across pages.
        inline_svg = (
            inline_svg
            .replace('id="nndv-figure-ir"', f'id="nndv-figure-ir-page-{index + 1}"')
            .replace('id="nndv-arrow"', f'id="nndv-arrow-page-{index + 1}"')
            .replace('url(#nndv-arrow)', f'url(#nndv-arrow-page-{index + 1})')
        )
        page_sections.append(
            f'<section class="page" data-page-index="{index}" data-page-id="{_esc(page_id)}" '
            f'data-width-mm="{page.width_mm:g}" data-height-mm="{page.height_mm:g}" '
            f'{"" if index == 0 else "hidden"}>{inline_svg}</section>'
        )
    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; img-src data:; style-src 'unsafe-inline'; script-src 'unsafe-inline'">
<meta name="generator" content="NN_DaVinci Figure Export {FIGURE_EXPORT_VERSION}">
<title>{_esc(figure.name)} — NN_DaVinci offline figure</title>
<style>
:root {{ color-scheme: light; font-family: system-ui,sans-serif; color:#0f172a; background:#e2e8f0 }}
body {{ margin:0 }}
header {{ position:sticky; top:0; z-index:3; display:flex; flex-wrap:wrap; gap:.5rem; align-items:center; padding:.65rem; background:#fff; border-bottom:1px solid #cbd5e1 }}
button,select,input {{ font:inherit }} button,select {{ padding:.3rem .55rem }}
main {{ padding:1rem; overflow:auto }}
.page {{ width:max-content; margin:auto; background:#fff; box-shadow:0 3px 18px #64748b66 }}
.page svg {{ display:block; max-width:calc(100vw - 2rem); height:auto }}
[data-figure-object-id] {{ cursor:pointer }} .selected {{ outline:1px dashed #e11d48; outline-offset:2px }}
#inspector {{ margin-left:auto; display:flex; gap:.4rem; align-items:center }}
#source {{ white-space:pre-wrap; max-height:45vh; overflow:auto; background:#0f172a; color:#e2e8f0; padding:1rem }}
.hint {{ font-size:.8rem; color:#475569 }}
</style>
</head>
<body data-nndv-document="offline-editable-figure" data-source-digest="{figure.digest()}">
<header>
<strong>{_esc(figure.name)}</strong>
<button id="previous" type="button" aria-label="Previous page">←</button>
<select id="page-select" aria-label="Page">{page_options}</select>
<button id="next" type="button" aria-label="Next page">→</button>
<button id="toggle-source" type="button">Source IR</button>
<button id="download-source" type="button">Download source</button>
<button id="download-svg" type="button">Download edited page SVG</button>
<button id="reset" type="button">Reset edits</button>
<span class="hint">Select and drag an object, or edit selected text below.</span>
<label id="inspector">Selected text <input id="text-value" disabled aria-label="Selected object text"></label>
</header>
<main>{''.join(page_sections)}<pre id="source" hidden></pre></main>
<script id="nndv-figure-ir" type="application/json">{source_json}</script>
<script id="nndv-export-policy" type="application/json">{policy_json}</script>
<script>
(() => {{
  'use strict';
  const sourceText = document.getElementById('nndv-figure-ir').textContent;
  const source = JSON.parse(sourceText);
  const pages = [...document.querySelectorAll('.page')];
  const select = document.getElementById('page-select');
  const textValue = document.getElementById('text-value');
  let current = 0, selectedId = '', drag = null;
  function activePage() {{ return pages[current]; }}
  function members() {{ return [...activePage().querySelectorAll('[data-figure-object-id]')].filter(el => el.dataset.figureObjectId === selectedId); }}
  function show(index) {{
    current = Math.max(0, Math.min(pages.length - 1, index));
    pages.forEach((page, i) => page.hidden = i !== current); select.value = String(current); clearSelection();
  }}
  function clearSelection() {{ pages.forEach(page => page.querySelectorAll('.selected').forEach(el => el.classList.remove('selected'))); selectedId=''; textValue.value=''; textValue.disabled=true; }}
  function selectObject(element) {{
    clearSelection(); selectedId = element.dataset.figureObjectId || '';
    members().forEach(el => el.classList.add('selected'));
    const text = members().find(el => el.tagName.toLowerCase() === 'text');
    textValue.disabled = !text; textValue.value = text ? text.textContent : '';
  }}
  pages.forEach(page => page.addEventListener('pointerdown', event => {{
    const element = event.target.closest('[data-figure-object-id]'); if (!element) return;
    selectObject(element); drag = {{x:event.clientX,y:event.clientY, transforms:members().map(el => el.getAttribute('transform') || '')}};
    element.setPointerCapture?.(event.pointerId); event.preventDefault();
  }}));
  addEventListener('pointermove', event => {{
    if (!drag || !selectedId) return;
    const svg = activePage().querySelector('svg'); const point = svg.createSVGPoint(); point.x=event.clientX; point.y=event.clientY;
    const now = point.matrixTransform(svg.getScreenCTM().inverse()); point.x=drag.x; point.y=drag.y;
    const then = point.matrixTransform(svg.getScreenCTM().inverse()); const dx=now.x-then.x, dy=now.y-then.y;
    members().forEach((el,i) => el.setAttribute('transform', `${{drag.transforms[i]}} translate(${{dx}} ${{dy}})`.trim()));
  }});
  addEventListener('pointerup', () => drag=null);
  textValue.addEventListener('input', () => members().filter(el => el.tagName.toLowerCase()==='text').forEach(el => {{
    const spans=[...el.querySelectorAll('tspan')]; if (spans.length) {{ spans[0].textContent=textValue.value; spans.slice(1).forEach(x=>x.remove()); }} else el.textContent=textValue.value;
  }}));
  select.addEventListener('change', () => show(Number(select.value)));
  document.getElementById('previous').onclick=()=>show(current-1); document.getElementById('next').onclick=()=>show(current+1);
  const sourcePanel=document.getElementById('source'); sourcePanel.textContent=JSON.stringify(source,null,2);
  document.getElementById('toggle-source').onclick=()=>sourcePanel.hidden=!sourcePanel.hidden;
  function download(name, type, text) {{ const a=document.createElement('a'); a.href=URL.createObjectURL(new Blob([text],{{type}})); a.download=name; a.click(); setTimeout(()=>URL.revokeObjectURL(a.href),0); }}
  document.getElementById('download-source').onclick=()=>download('figure.nndv.json','application/json',JSON.stringify(source,null,2)+'\\n');
  document.getElementById('download-svg').onclick=()=>{{
    const pageNumber=current+1;
    const serialized=new XMLSerializer().serializeToString(activePage().querySelector('svg'))
      .replace(`nndv-figure-ir-page-${{pageNumber}}`,'nndv-figure-ir')
      .replaceAll(`nndv-arrow-page-${{pageNumber}}`,'nndv-arrow');
    download(`figure.page-${{String(pageNumber).padStart(3,'0')}}.svg`,'image/svg+xml',serialized+'\\n');
  }};
  document.getElementById('reset').onclick=()=>location.reload(); show(0);
}})();
</script>
</body>
</html>
"""


def export_figure_html(figure: FigureIR, destination: str | Path) -> Path:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(render_figure_html(figure), encoding="utf-8")
    return destination


def _pptx_color(value: str) -> Any:
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(_hex(value).lstrip("#").upper())


def _pptx_polygon(slide: Any, points: list[list[float]], *, name: str) -> Any:
    coords = [[float(x) * MM_TO_EMU, float(y) * MM_TO_EMU] for x, y in points]
    builder = slide.shapes.build_freeform(coords[0][0], coords[0][1])
    builder.add_line_segments([(x, y) for x, y in coords[1:]], close=True)
    shape = builder.convert_to_shape()
    shape.name = name
    return shape


def _primitive_values_on_pptx_slide(
    primitive: CompiledPrimitive,
    *,
    compiled_page_offset: float,
    slide_x_offset: float,
    slide_y_offset: float,
) -> dict[str, Any]:
    values = copy.deepcopy(primitive.values)
    for key in ("x", "x1", "x2", "cx"):
        if key in values:
            values[key] = float(values[key]) + slide_x_offset
    for key in ("y", "y1", "y2", "cy"):
        if key in values:
            values[key] = float(values[key]) - compiled_page_offset + slide_y_offset
    if "points" in values:
        values["points"] = [
            [float(x) + slide_x_offset, float(y) - compiled_page_offset + slide_y_offset]
            for x, y in values["points"]
        ]
    for line in values.get("lines", []):
        line["x"] = float(line["x"]) + slide_x_offset
        line["baseline_y"] = float(line["baseline_y"]) - compiled_page_offset + slide_y_offset
    for box_name in ("text_box", "measured_bbox"):
        if box_name in values:
            values[box_name]["x"] = float(values[box_name]["x"]) + slide_x_offset
            values[box_name]["y"] = float(values[box_name]["y"]) - compiled_page_offset + slide_y_offset
    return values


def export_figure_pptx(figure: FigureIR, destination: str | Path) -> Path:
    """Write one editable slide per page using a deterministic 1:1 canvas.

    PPTX has a presentation-wide slide size.  For heterogeneous pages, the
    slide canvas is the maximum source width and maximum source height.  Each
    FigurePage is centered on that canvas without scaling or cropping.
    """

    figure.validate()
    try:
        from pptx import Presentation
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Emu, Mm, Pt
    except ImportError as exc:
        raise OptionalDependencyError("Editable Figure IR PowerPoint export requires python-pptx") from exc
    presentation = Presentation()
    pages = _ordered_pages(figure)
    canvas_width = max(page.width_mm for page in pages)
    canvas_height = max(page.height_mm for page in pages)
    presentation.slide_width = Mm(canvas_width)
    presentation.slide_height = Mm(canvas_height)
    presentation.core_properties.title = figure.name
    presentation.core_properties.subject = "Editable NN_DaVinci FigureIR export"
    # OOXML core-property text is capped at 255 characters by python-pptx.
    # Full page identity/order remains in the ordered slide shape names and in
    # ``figure_export_policy``; keep only the stable audit marker here.
    presentation.core_properties.comments = (
        f"NNDV-PPTX-1;digest={figure.digest()};pages={len(pages)};"
        f"canvas_mm={canvas_width:g}x{canvas_height:g};policy=max-centered-1:1-no-crop"
    )
    blank = presentation.slide_layouts[6]
    by_page: dict[str, list[tuple[str, str, CompiledPrimitive]]] = {}
    for page_id, panel_id, layer_id, primitive in compile_figure(figure):
        by_page.setdefault(page_id, []).append((panel_id, layer_id, primitive))
    page_offset = 0.0
    for page_index, page in enumerate(pages):
        slide = presentation.slides.add_slide(blank)
        slide_x_offset = (canvas_width - page.width_mm) / 2.0
        slide_y_offset = (canvas_height - page.height_mm) / 2.0
        page_frame = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Mm(slide_x_offset),
            Mm(slide_y_offset),
            Mm(page.width_mm),
            Mm(page.height_mm),
        )
        page_frame.name = f"NNDV Page {page_index + 1:03d} {page.id} {page.title}"
        page_frame.fill.solid()
        page_frame.fill.fore_color.rgb = _pptx_color(figure.design_tokens["color.background"])
        page_frame.line.fill.background()
        for panel in sorted(page.panels, key=lambda item: (item.order, item.id)):
            box = panel.geometry
            boundary = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Mm(box["x"] + slide_x_offset),
                Mm(box["y"] + slide_y_offset),
                Mm(box["width"]),
                Mm(box["height"]),
            )
            boundary.name = f"NNDV Panel {panel.label} {panel.id}"
            boundary.fill.background()
            boundary.line.color.rgb = _pptx_color("#cbd5e1")
            boundary.line.width = Pt(0.35)
            panel_font_pt = max(7.0, float(figure.design_tokens["font.size.pt"]))
            label = slide.shapes.add_textbox(
                Mm(box["x"] + slide_x_offset + 1.5),
                Mm(box["y"] + slide_y_offset + 0.8),
                Mm(12.0),
                Mm(pt_to_mm(panel_font_pt * 1.4)),
            )
            label.name = f"NNDV editable panel label {panel.id}"
            label.text_frame.clear()
            label.text_frame.margin_left = label.text_frame.margin_right = Emu(0)
            label.text_frame.margin_top = label.text_frame.margin_bottom = Emu(0)
            label.text_frame.paragraphs[0].text = panel.label
            label.text_frame.paragraphs[0].font.name = str(figure.design_tokens["font.family"]).split(",")[0]
            label.text_frame.paragraphs[0].font.size = Pt(panel_font_pt)
            label.text_frame.paragraphs[0].font.bold = True
            label.text_frame.paragraphs[0].font.color.rgb = _pptx_color(figure.design_tokens["color.foreground"])
        for _panel_id, _layer_id, primitive in by_page.get(page.id, []):
            values = _primitive_values_on_pptx_slide(
                primitive,
                compiled_page_offset=page_offset,
                slide_x_offset=slide_x_offset,
                slide_y_offset=slide_y_offset,
            )
            style = primitive.style
            shape = None
            if primitive.kind == "rect":
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Mm(values["x"]), Mm(values["y"]), Mm(values["width"]), Mm(values["height"]))
            elif primitive.kind == "polygon":
                shape_name = (
                    f"NNDV tensor face {values['face']} {primitive.object_id}"
                    if values.get("face")
                    else f"NNDV editable polygon {primitive.object_id}"
                )
                shape = _pptx_polygon(slide, values["points"], name=shape_name)
            elif primitive.kind in {"polyline", "line"}:
                points = values.get("points") or [[values["x1"], values["y1"]], [values["x2"], values["y2"]]]
                segments = list(zip(points, points[1:]))
                for index, (start, end) in enumerate(segments):
                    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Mm(start[0]), Mm(start[1]), Mm(end[0]), Mm(end[1]))
                    connector.name = f"NNDV editable connector {primitive.semantic} {_panel_id} {_layer_id} {primitive.object_id} {index}"
                    connector.line.color.rgb = _pptx_color(style["stroke"])
                    connector.line.width = Pt(style["stroke_width"])
                    if style.get("dash"):
                        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
                    line_xml = connector.line._get_or_add_ln()
                    if index == 0 and values.get("arrow") == "both":
                        head = OxmlElement("a:headEnd")
                        head.set("type", "triangle")
                        line_xml.append(head)
                    if index == len(segments) - 1 and values.get("arrow") in {"end", "both", "true"}:
                        tail = OxmlElement("a:tailEnd")
                        tail.set("type", "triangle")
                        line_xml.append(tail)
                continue
            elif primitive.kind == "circle":
                shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Mm(values["cx"] - values["r"]), Mm(values["cy"] - values["r"]), Mm(2 * values["r"]), Mm(2 * values["r"]))
            elif primitive.kind == "text":
                text_box = values.get("text_box") or {
                    "x": values["x"], "y": float(values["y"]) - 3.5, "width": 45.0, "height": 8.0,
                }
                shape = slide.shapes.add_textbox(
                    Mm(float(text_box["x"])),
                    Mm(float(text_box["y"])),
                    Mm(float(text_box["width"])),
                    Mm(float(text_box["height"])),
                )
                frame = shape.text_frame
                frame.clear()
                frame.margin_left = frame.margin_right = Emu(0)
                frame.margin_top = frame.margin_bottom = Emu(0)
                frame.word_wrap = False
                frame.vertical_anchor = {
                    "top": MSO_ANCHOR.TOP,
                    "bottom": MSO_ANCHOR.BOTTOM,
                }.get(str(values.get("vertical_anchor", "middle")), MSO_ANCHOR.MIDDLE)
                paragraph = frame.paragraphs[0]
                paragraph.text = "\n".join(str(line["text"]) for line in values.get("lines", [])) or str(values["text"])
                paragraph.alignment = {
                    "middle": PP_ALIGN.CENTER,
                    "end": PP_ALIGN.RIGHT,
                }.get(str(values.get("anchor", "start")), PP_ALIGN.LEFT)
                paragraph.font.name = str(style["font_family"]).split(",")[0]
                paragraph.font.size = Pt(float(style["font_size_pt"]))
                paragraph.font.bold = int(style["font_weight"]) >= 600
                paragraph.font.color.rgb = _pptx_color(figure.design_tokens["color.foreground"])
                paragraph.line_spacing = Pt(float(values.get("line_height_pt", style["line_height_pt"])))
                if "formula_source" in values:
                    formula_record = {
                        key: values[key]
                        for key in (
                            "formula_source", "formula_syntax_version", "formula_valid",
                            "formula_fallback", "formula_error", "formula_features",
                        )
                    }
                    shape.name = f"NNDV editable formula {primitive.object_id}"
                    # ``descr`` is inert OOXML alternative text.  It retains
                    # the complete normalization decision without converting
                    # the equation into a picture or executable TeX object.
                    shape._element.nvSpPr.cNvPr.set(
                        "descr",
                        json.dumps(formula_record, ensure_ascii=False, sort_keys=True, separators=(",", ":")),
                    )
                else:
                    shape.name = f"NNDV editable text {primitive.object_id}"
                continue
            elif primitive.kind == "image" and str(values.get("href", "")).startswith("data:"):
                try:
                    encoded = str(values["href"]).split(",", 1)[1]
                    shape = slide.shapes.add_picture(io.BytesIO(base64.b64decode(encoded)), Mm(values["x"]), Mm(values["y"]), Mm(values["width"]), Mm(values["height"]))
                except Exception as exc:
                    raise ExportError(f"Editable PPTX image object {primitive.object_id!r} is malformed: {exc}") from exc
            elif primitive.kind == "image":
                raise ExportError(
                    f"Editable PPTX image object {primitive.object_id!r} must use an embedded data URI"
                )
            if shape is not None:
                shape.name = shape.name if shape.name.startswith("NNDV tensor") else f"NNDV editable {primitive.kind} {primitive.object_id}"
                if hasattr(shape, "fill"):
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = _pptx_color(values.get("fill", style["fill"]))
                    shape.fill.transparency = int(round((1.0 - style["opacity"]) * 100))
                if hasattr(shape, "line"):
                    shape.line.color.rgb = _pptx_color(style["stroke"])
                    shape.line.width = Pt(style["stroke_width"])
        page_offset += page.height_mm + 6.0
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(destination)
    return destination


Box = tuple[float, float, float, float]


def _primitive_bbox(primitive: CompiledPrimitive) -> Box | None:
    """Return the compiled primitive's conservative local-page bounds."""

    values = primitive.values
    if primitive.kind in {"rect", "image"}:
        x, y = float(values["x"]), float(values["y"])
        return x, y, x + float(values["width"]), y + float(values["height"])
    if primitive.kind == "circle":
        cx, cy, radius = float(values["cx"]), float(values["cy"]), float(values["r"])
        return cx - radius, cy - radius, cx + radius, cy + radius
    if primitive.kind == "text" and values.get("measured_bbox"):
        box = values["measured_bbox"]
        x, y = float(box["x"]), float(box["y"])
        return x, y, x + float(box["width"]), y + float(box["height"])
    if primitive.kind in {"polygon", "polyline", "line"}:
        points = values.get("points")
        if points is None and primitive.kind == "line":
            points = [[values["x1"], values["y1"]], [values["x2"], values["y2"]]]
        if points:
            xs = [float(point[0]) for point in points]
            ys = [float(point[1]) for point in points]
            return min(xs), min(ys), max(xs), max(ys)
    return None


def _union_boxes(boxes: Iterable[Box]) -> Box | None:
    values = list(boxes)
    if not values:
        return None
    return (
        min(box[0] for box in values),
        min(box[1] for box in values),
        max(box[2] for box in values),
        max(box[3] for box in values),
    )


def _overlaps(first: Box, second: Box, tolerance: float = 0.05) -> bool:
    return first[0] < second[2] - tolerance and first[2] > second[0] + tolerance and first[1] < second[3] - tolerance and first[3] > second[1] + tolerance


def _segment_intersects_box(start: list[float], end: list[float], box: Box, tolerance: float = 0.05) -> bool:
    """Liang–Barsky segment/rectangle test, including crossings between vertices."""

    left, top, right, bottom = (
        box[0] + tolerance,
        box[1] + tolerance,
        box[2] - tolerance,
        box[3] - tolerance,
    )
    if left >= right or top >= bottom:
        return False
    x0, y0 = float(start[0]), float(start[1])
    dx, dy = float(end[0]) - x0, float(end[1]) - y0
    lower, upper = 0.0, 1.0
    for p, q in ((-dx, x0 - left), (dx, right - x0), (-dy, y0 - top), (dy, bottom - y0)):
        if abs(p) < 1e-12:
            if q < 0:
                return False
            continue
        ratio = q / p
        if p < 0:
            lower = max(lower, ratio)
        else:
            upper = min(upper, ratio)
        if lower > upper:
            return False
    return True


def _segment_crosses(
    first_start: list[float], first_end: list[float], second_start: list[float], second_end: list[float]
) -> bool:
    def orientation(a: list[float], b: list[float], c: list[float]) -> float:
        return (b[0] - a[0]) * (c[1] - a[1]) - (b[1] - a[1]) * (c[0] - a[0])

    values = (
        orientation(first_start, first_end, second_start),
        orientation(first_start, first_end, second_end),
        orientation(second_start, second_end, first_start),
        orientation(second_start, second_end, first_end),
    )
    # Collinear/endpoint contacts are handled by explicit merge provenance,
    # not counted as unexplained crossings here.
    if any(abs(value) < 1e-8 for value in values):
        return False
    return (values[0] > 0) != (values[1] > 0) and (values[2] > 0) != (values[3] > 0)


def figure_proof(figure: FigureIR) -> dict[str, Any]:
    """Run a conservative compiled-IR preflight.

    This report deliberately is *not* a final-output geometry oracle.  Browser
    DOM/CTM measurements of the serialized SVG are required for release-blocking
    text scale, clipping, overlap, stroke and marker claims.
    """

    figure.validate()
    overlaps: list[list[str]] = []
    text_conflicts: list[list[str]] = []
    text_graphic_conflicts: list[list[str]] = []
    clipping: list[str] = []
    page_clipping: list[str] = []
    edge_collisions: list[list[str]] = []
    crossings: list[list[str]] = []
    marker_clipping: list[str] = []
    object_count = 0
    provenance_incomplete: list[str] = []
    font_sizes: list[float] = []
    stroke_widths: list[float] = []
    transparent_objects: list[str] = []
    for page in figure.pages:
        page_box: Box = (0.0, 0.0, float(page.width_mm), float(page.height_mm))
        for panel in page.panels:
            objects = [item for layer in panel.layers if layer.visible for item in layer.objects if item.visible]
            object_count += len(objects)
            panel_box: Box = (
                float(panel.geometry["x"]),
                float(panel.geometry["y"]),
                float(panel.geometry["x"] + panel.geometry["width"]),
                float(panel.geometry["y"] + panel.geometry["height"]),
            )
            compiled_by_object: dict[str, list[CompiledPrimitive]] = {}
            object_by_id = {item.id: item for item in objects}
            for item in objects:
                primitives = compile_object(figure, item, panel_geometry=panel.geometry)
                compiled_by_object[item.id] = primitives
                for primitive in primitives:
                    if primitive.kind == "text":
                        font_sizes.append(float(primitive.style["font_size_pt"]))
                    elif primitive.kind not in {"image"}:
                        stroke_widths.append(float(primitive.style["stroke_width_pt"]))
                    if float(primitive.style.get("opacity", 1.0)) <= 0:
                        transparent_objects.append(item.id)

            graphic_boxes: dict[str, Box] = {}
            text_boxes: list[tuple[str, Box]] = []
            all_boxes: list[tuple[str, Box]] = []
            for object_id, primitives in compiled_by_object.items():
                graphic = _union_boxes(
                    box
                    for primitive in primitives
                    if primitive.kind in {"rect", "image", "circle", "polygon"}
                    if (box := _primitive_bbox(primitive)) is not None
                )
                if graphic is not None:
                    graphic_boxes[object_id] = graphic
                    all_boxes.append((object_id, graphic))
                for primitive in primitives:
                    if primitive.kind != "text":
                        continue
                    text_box = _primitive_bbox(primitive)
                    if text_box is not None:
                        text_boxes.append((object_id, text_box))
                        all_boxes.append((object_id, text_box))

            panel_font_pt = max(7.0, float(figure.design_tokens["font.size.pt"]))
            panel_measurement = measure_text_line(
                panel.label,
                font_size_pt=panel_font_pt,
                font_family=str(figure.design_tokens["font.family"]),
                font_weight=700,
            )
            font_sizes.append(panel_font_pt)
            panel_label_id = f"panel-label:{panel.id}"
            panel_label_box: Box = (
                panel_box[0] + 1.5,
                panel_box[1] + 1.2,
                panel_box[0] + 1.5 + panel_measurement.width_mm,
                panel_box[1] + 1.2 + panel_measurement.height_mm,
            )
            text_boxes.append((panel_label_id, panel_label_box))
            all_boxes.append((panel_label_id, panel_label_box))
            stroke_widths.append(0.35)

            for object_id, box in all_boxes:
                if box[0] < panel_box[0] - 1e-6 or box[1] < panel_box[1] - 1e-6 or box[2] > panel_box[2] + 1e-6 or box[3] > panel_box[3] + 1e-6:
                    clipping.append(object_id)
                if box[0] < page_box[0] - 1e-6 or box[1] < page_box[1] - 1e-6 or box[2] > page_box[2] + 1e-6 or box[3] > page_box[3] + 1e-6:
                    page_clipping.append(object_id)

            content_boxes = [
                (object_id, box)
                for object_id, box in graphic_boxes.items()
                if object_by_id[object_id].kind in {"node-glyph", "operator-glyph", "tensor-glyph", "image", "external-vector-group"}
            ]
            for index, (object_id, box) in enumerate(content_boxes):
                for other_id, other_box in content_boxes[index + 1:]:
                    if _overlaps(box, other_box):
                        overlaps.append([object_id, other_id])

            for index, (object_id, box) in enumerate(text_boxes):
                for other_id, other_box in text_boxes[index + 1:]:
                    if object_id != other_id and _overlaps(box, other_box):
                        text_conflicts.append([object_id, other_id])
                for graphic_id, graphic_box in graphic_boxes.items():
                    if graphic_id != object_id and _overlaps(box, graphic_box):
                        text_graphic_conflicts.append([object_id, graphic_id])
            for edge in (item for item in objects if item.kind == "edge"):
                route = edge.manual_route
                if not route:
                    continue
                for node_id, box in graphic_boxes.items():
                    if node_id in edge.metadata.get("endpoint_object_ids", []):
                        continue
                    if any(_segment_intersects_box(start, end, box) for start, end in zip(route, route[1:])):
                        edge_collisions.append([edge.id, node_id])
                if edge.style.resolved(figure.design_tokens).get("arrow", figure.design_tokens["arrow.kind"]) in {"end", "both", "true"}:
                    arrow_mm = pt_to_mm(float(edge.style.resolved(figure.design_tokens).get("arrow_size", figure.design_tokens["arrow.size.pt"])))
                    end_x, end_y = float(route[-1][0]), float(route[-1][1])
                    if end_x - arrow_mm < panel_box[0] or end_y - arrow_mm < panel_box[1] or end_x + arrow_mm > panel_box[2] or end_y + arrow_mm > panel_box[3]:
                        marker_clipping.append(edge.id)
            routed_edges = [item for item in objects if item.kind == "edge" and len(item.manual_route) >= 2]
            for index, edge in enumerate(routed_edges):
                for other in routed_edges[index + 1:]:
                    shared_endpoints = set(edge.metadata.get("endpoint_object_ids", [])).intersection(
                        other.metadata.get("endpoint_object_ids", [])
                    )
                    if shared_endpoints or edge.metadata.get("crossing_explained") or other.metadata.get("crossing_explained"):
                        continue
                    if any(
                        _segment_crosses(first, second, other_first, other_second)
                        for first, second in zip(edge.manual_route, edge.manual_route[1:])
                        for other_first, other_second in zip(other.manual_route, other.manual_route[1:])
                    ):
                        crossings.append([edge.id, other.id])
            for item in objects:
                if item.provenance.kind == "unknown" and not item.provenance.reason:
                    provenance_incomplete.append(item.id)
    minimum_font = min(font_sizes) if font_sizes else None
    minimum_stroke = min(stroke_widths) if stroke_widths else None
    unique_overlaps = sorted({tuple(pair) for pair in overlaps})
    unique_text_conflicts = sorted({tuple(pair) for pair in text_conflicts})
    unique_text_graphic_conflicts = sorted({tuple(pair) for pair in text_graphic_conflicts})
    clipping = sorted(set(clipping))
    page_clipping = sorted(set(page_clipping))
    unique_edge_collisions = sorted({tuple(pair) for pair in edge_collisions})
    unique_crossings = sorted({tuple(pair) for pair in crossings})
    marker_clipping = sorted(set(marker_clipping))
    transparent_objects = sorted(set(transparent_objects))
    background = _hex(figure.design_tokens["color.background"], "#ffffff")
    contrast = {
        "text": round(contrast_ratio(_hex(figure.design_tokens["color.foreground"], "#0f172a"), background), 3),
        "data_edge": round(contrast_ratio(_hex(figure.design_tokens["line.data.color"], "#334155"), background), 3),
        "annotation_edge": round(contrast_ratio(_hex(figure.design_tokens["line.annotation.color"], "#7c3aed"), background), 3),
    }
    tensor_luminance = [
        relative_luminance(_hex(figure.design_tokens[key]))
        for key in ("color.tensor.front", "color.tensor.top", "color.tensor.side")
    ]
    grayscale_distinguishable = len({round(value, 2) for value in tensor_luminance}) >= 2
    report = {
        "schema_version": "nndv-figure-proof-2",
        "figure_ir_version": figure.schema_version,
        "figure_digest": figure.digest(),
        "measurement_source": "compiled Figure IR estimate; strict final-SVG Chrome DOM/CTM oracle required",
        "final_output_verified": False,
        "release_blocker_eligible": False,
        "page_count": len(figure.pages),
        "panel_count": sum(len(page.panels) for page in figure.pages),
        "object_count": object_count,
        "overflow": 0 if not clipping else len(clipping),
        "node_tensor_overlap": len(unique_overlaps),
        "overlap_pairs": [list(pair) for pair in unique_overlaps],
        "text_conflict": len(unique_text_conflicts),
        "text_conflict_pairs": [list(pair) for pair in unique_text_conflicts],
        "text_graphic_conflict": len(unique_text_graphic_conflicts),
        "text_graphic_conflict_pairs": [list(pair) for pair in unique_text_graphic_conflicts],
        "edge_node_collision": len(unique_edge_collisions),
        "edge_collision_pairs": [list(pair) for pair in unique_edge_collisions],
        "clipping": len(clipping),
        "clipped_object_ids": clipping,
        "page_clipping": len(page_clipping),
        "page_clipped_object_ids": page_clipping,
        "marker_clipping": len(marker_clipping),
        "marker_clipped_edge_ids": marker_clipping,
        "unexplained_crossing": len(unique_crossings),
        "unexplained_crossing_pairs": [list(pair) for pair in unique_crossings],
        "minimum_font_pt": minimum_font,
        "minimum_stroke_pt": minimum_stroke,
        "horizontal_text_scale": None,
        "transform_text_scale": None,
        "transparent_object_ids": transparent_objects,
        "provenance_incomplete": provenance_incomplete,
        "contrast_ratios": contrast,
        "contrast_pass": contrast["text"] >= 4.5 and contrast["data_edge"] >= 3.0 and contrast["annotation_edge"] >= 3.0,
        "grayscale_distinguishable": grayscale_distinguishable,
        "annotation_data_edge_distinction": "dash + semantic metadata + colour",
        "provisional_pass": not overlaps and not text_conflicts and not text_graphic_conflicts and not clipping and not page_clipping and not edge_collisions and not crossings and not marker_clipping and minimum_font is not None and minimum_font >= 7.0 and minimum_stroke is not None and minimum_stroke >= 0.1 and not provenance_incomplete and not transparent_objects and contrast["text"] >= 4.5 and contrast["data_edge"] >= 3.0 and contrast["annotation_edge"] >= 3.0 and grayscale_distinguishable,
    }
    report["pass"] = report["provisional_pass"]
    return report


def provenance_report(figure: FigureIR) -> dict[str, Any]:
    records = []
    for item in figure.iter_objects():
        records.append({"object_id": item.id, "object_kind": item.kind, **item.provenance.__dict__} if hasattr(item.provenance, "__dict__") else {"object_id": item.id, "object_kind": item.kind, "kind": item.provenance.kind, "source_id": item.provenance.source_id, "graph_ir_ids": item.provenance.graph_ir_ids, "source_locator": item.provenance.source_locator, "evidence": item.provenance.evidence, "reason": item.provenance.reason, "author_annotation": item.provenance.author_annotation})
    return {
        "schema_version": "nndv-figure-provenance-1", "figure_digest": figure.digest(),
        "records": sorted(records, key=lambda item: item["object_id"]),
        "author_annotations_are_graph_ir": False,
    }


def export_figure(
    figure: FigureIR,
    output: str | Path,
    *,
    formats: Iterable[str] = ("svg",),
) -> list[Path]:
    """Export requested formats from a single validated FigureIR.

    A one-page figure retains all legacy filenames.  Multi-page formats that
    cannot represent heterogeneous physical page sizes emit ordered
    ``.page-001`` files; PDF, PPTX, and HTML remain single containers.  The
    returned list includes every emitted path in request order and then page
    order.
    """

    figure.validate()
    stem = Path(output).with_suffix("")
    stem.parent.mkdir(parents=True, exist_ok=True)
    page_count = len(figure.pages)
    outputs: list[Path] = []
    for requested in formats:
        name = requested.lower().lstrip(".")
        if name == "svg":
            for index, (_page_id, svg) in enumerate(render_figure_svg_pages(figure)):
                destination = _page_destination(stem, ".svg", index, page_count)
                destination.write_text(svg, encoding="utf-8")
                outputs.append(destination)
        elif name == "pdf":
            destination = stem.with_suffix(".pdf")
            outputs.append(export_figure_pdf(figure, destination))
        elif name in {"tikz", "tex"}:
            for index, (_page_id, tikz) in enumerate(render_figure_tikz_pages(figure)):
                destination = _page_destination(stem, ".tex", index, page_count)
                destination.write_text(tikz, encoding="utf-8")
                outputs.append(destination)
        elif name in {"pptx", "powerpoint"}:
            destination = stem.with_suffix(".pptx")
            outputs.append(export_figure_pptx(figure, destination))
        elif name == "png":
            outputs.extend(export_figure_png_pages(figure, stem))
        elif name in {"eps", "epsf"}:
            outputs.extend(export_figure_eps_pages(figure, stem))
        elif name in {"html", "htm"}:
            destination = stem.with_suffix(".html")
            outputs.append(export_figure_html(figure, destination))
        else:
            raise ExportError(
                "Figure IR supports SVG, PDF, TikZ/TeX, editable PPTX, 300-DPI PNG, vector EPS, and offline HTML"
            )
    return outputs


def export_submission_package(
    figure: FigureIR,
    directory: str | Path,
    *,
    project: dict[str, Any] | None = None,
    caption: str | None = None,
) -> list[Path]:
    destination = Path(directory)
    destination.mkdir(parents=True, exist_ok=True)
    outputs = export_figure(
        figure,
        destination / "figure",
        formats=("svg", "pdf", "tikz", "pptx", "png", "eps", "html"),
    )
    project_document = project or {
        "project_version": "1.4", "name": figure.name,
        "graph": {"name": f"{figure.name} (no model graph attached)", "nodes": [], "edges": [], "subgraphs": [], "annotations": [], "constraints": [], "inputs": [], "outputs": [], "metadata": {"figure_only": True, "model_facts_claimed": False}, "analysis": {}, "ir_version": "1.0"},
        "figure_ir": figure.to_dict(),
        "scene_ir": empty_scene(f"{figure.name} · 3D").to_dict(),
    }
    nndv = destination / "figure.nndv.json"
    nndv.write_text(json.dumps(project_document, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    captions = [str(item.metadata.get("text", item.name)) for item in figure.iter_objects() if item.kind == "caption"]
    caption_path = destination / "caption.md"
    caption_path.write_text((caption if caption is not None else "\n\n".join(captions) or f"Figure: {figure.name}") + "\n", encoding="utf-8")
    provenance_path = destination / "provenance.json"
    provenance_path.write_text(json.dumps(provenance_report(figure), ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    proof_path = destination / "proof.json"
    proof_path.write_text(json.dumps(figure_proof(figure), ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    policy_path = destination / "export-policy.json"
    policy_path.write_text(
        json.dumps(figure_export_policy(figure), ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return [*outputs, nndv, caption_path, provenance_path, proof_path, policy_path]


__all__ = [
    "FIGURE_EXPORT_VERSION", "FIGURE_SVG_METADATA_VERSION", "PNG_EXPORT_DPI", "CompiledPrimitive", "compile_figure",
    "compile_object", "export_figure", "export_figure_eps_pages", "export_figure_html", "export_figure_pdf",
    "export_figure_png_pages", "export_figure_pptx", "export_submission_package", "figure_export_policy",
    "figure_proof", "import_figure_svg", "provenance_report", "render_figure_html", "render_figure_svg",
    "render_figure_svg_pages", "render_figure_tikz", "render_figure_tikz_pages",
]
