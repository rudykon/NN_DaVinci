from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from ..errors import ExportError, OptionalDependencyError
from ..ir import GraphIR
from ..layout import LayoutResult
from .svg import SvgRenderer
from .tikz import TikzRenderer
from .html import HtmlRenderer

CUSTOM_EXPORTERS: dict[str, Any] = {}


def register_exporter(name: str, exporter: Any, *, replace: bool = False) -> None:
    normalized = name.lower().lstrip(".")
    if normalized in CUSTOM_EXPORTERS and not replace:
        raise ValueError(f"Exporter {normalized!r} is already registered")
    CUSTOM_EXPORTERS[normalized] = exporter


def export_graph(
    graph: GraphIR,
    layout: LayoutResult,
    theme: dict[str, Any],
    output: str | Path,
    *,
    formats: Iterable[str] | None = None,
    **options: Any,
) -> list[Path]:
    target = Path(output)
    requested = [item.lower().lstrip(".") for item in formats] if formats else [target.suffix.lower().lstrip(".") or "svg"]
    stem = target.with_suffix("") if target.suffix else target
    svg = SvgRenderer().render(graph, layout, theme, **options)
    outputs: list[Path] = []
    for format_name in requested:
        destination = stem.with_suffix(f".{format_name}")
        destination.parent.mkdir(parents=True, exist_ok=True)
        if format_name == "svg":
            destination.write_text(svg, encoding="utf-8")
        elif format_name in {"pdf", "png", "eps"}:
            try:
                import cairosvg
            except ImportError as exc:
                raise OptionalDependencyError(
                    f"{format_name.upper()} export requires CairoSVG",
                    hint="Install nn-davinci[export], or export SVG/TikZ.",
                ) from exc
            function = getattr(cairosvg, f"svg2{format_name}")
            kwargs: dict[str, Any] = {"bytestring": svg.encode("utf-8"), "write_to": str(destination)}
            if format_name == "png" and options.get("scale"):
                kwargs["scale"] = float(options["scale"])
            function(**kwargs)
        elif format_name in {"tikz", "tex"}:
            destination = stem.with_suffix(".tex")
            destination.write_text(TikzRenderer().render(graph, layout, theme, **options), encoding="utf-8")
        elif format_name in {"pptx", "powerpoint"}:
            destination = stem.with_suffix(".pptx")
            _export_pptx(graph, layout, theme, destination)
        elif format_name in {"html", "interactive"}:
            destination = stem.with_suffix(".html")
            destination.write_text(HtmlRenderer().render(graph, layout, theme, **options), encoding="utf-8")
        elif format_name in CUSTOM_EXPORTERS:
            result = CUSTOM_EXPORTERS[format_name](graph, layout, theme, destination, **options)
            if result is not None:
                destination = Path(result)
        else:
            raise ExportError(
                f"Unsupported export format {format_name!r}",
                hint="Use svg, pdf, tikz/tex, png, eps, pptx, or html.",
            )
        outputs.append(destination)
    return outputs


def _export_pptx(graph: GraphIR, layout: LayoutResult, theme: dict[str, Any], destination: Path) -> None:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise OptionalDependencyError(
            "Editable PowerPoint export requires python-pptx",
            hint="Install nn-davinci[export].",
        ) from exc
    presentation = Presentation()
    paper = layout.metadata.get("paper", {})
    physical_width_mm = paper.get("physical_width_mm")
    physical_height_mm = paper.get("physical_height_mm")
    if physical_width_mm and physical_height_mm:
        presentation.slide_width = Inches(float(physical_width_mm) / 25.4)
        presentation.slide_height = Inches(float(physical_height_mm) / 25.4)
    else:
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide_width_in = float(presentation.slide_width) / 914400
    slide_height_in = float(presentation.slide_height) / 914400
    scale = min((slide_width_in - 0.3) / max(layout.width, 1), (slide_height_in - 0.3) / max(layout.height, 1))
    x_offset, y_offset = 0.15, 0.15
    for route in layout.edges.values():
        for start, end in zip(route.points, route.points[1:]):
            connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x_offset + start[0] * scale), Inches(y_offset + start[1] * scale), Inches(x_offset + end[0] * scale), Inches(y_offset + end[1] * scale))
            connector.line.color.rgb = RGBColor.from_string(theme["edge"].lstrip("#"))
            connector.line.width = Pt(theme["edge_width"])
    nodes = graph.node_map()
    for node_id, item in layout.nodes.items():
        node = nodes.get(node_id)
        if node is None:
            continue
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_offset + item.x * scale), Inches(y_offset + item.y * scale), Inches(item.width * scale), Inches(item.height * scale))
        color = theme["category_colors"].get(node.category, theme["node_fill"])
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color.lstrip("#"))
        shape.line.color.rgb = RGBColor.from_string(theme["border"].lstrip("#"))
        frame = shape.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        paragraph.text = f"{node.name}\n{node.op_type}"
        paragraph.font.name = theme["font_family"].split(",")[0]
        paragraph.font.size = Pt(max(7, theme["font_size"] - 1))
    presentation.save(destination)
