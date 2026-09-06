from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from nn_davinci.adapters.jax import JaxAdapter, MlirAdapter
from nn_davinci.adapters.manual import ManualAdapter
from nn_davinci.layout import LayoutEngine
from nn_davinci.project import Project
from nn_davinci.render import HtmlRenderer, export_graph
from nn_davinci.themes import get_theme


ROOT = Path(__file__).parents[1]


class ExtendedIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.graph = ManualAdapter().load(ROOT / "examples" / "resnet.json")

    @unittest.skipUnless(importlib.util.find_spec("keras"), "Keras is optional")
    def test_keras_functional_graph(self):
        import keras
        from nn_davinci.adapters.keras import KerasAdapter

        inputs = keras.Input(shape=(8,), name="features")
        hidden = keras.layers.Dense(4, activation="relu", name="hidden")(inputs)
        outputs = keras.layers.Dense(2, name="scores")(hidden)
        graph = KerasAdapter().load(keras.Model(inputs, outputs, name="keras_tiny")).validate()
        self.assertEqual(len(graph.nodes), 3)
        self.assertEqual(len(graph.edges), 2)
        self.assertEqual(sum(node.parameters for node in graph.nodes), 46)

    @unittest.skipUnless(importlib.util.find_spec("jax"), "JAX is optional")
    def test_jaxpr_graph(self):
        import jax.numpy as jnp

        def function(value):
            return jnp.tanh(value @ jnp.ones((4, 3)))

        graph = JaxAdapter().load(function, sample_input=jnp.ones((2, 4))).validate()
        self.assertEqual(graph.metadata["source_format"], "jaxpr")
        self.assertTrue(any(node.op_type == "dot_general" for node in graph.nodes))
        self.assertGreaterEqual(len(graph.edges), 3)

    def test_mlir_ssa_graph(self):
        source = '''
          %0 = "stablehlo.constant"() : () -> tensor<2x2xf32>
          %1 = "stablehlo.add"(%0, %0) : (tensor<2x2xf32>, tensor<2x2xf32>) -> tensor<2x2xf32>
        '''
        graph = MlirAdapter().load(source).validate()
        self.assertEqual(len(graph.nodes), 2)
        self.assertEqual(len(graph.edges), 2)

    @unittest.skipUnless(importlib.util.find_spec("tensorflow"), "TensorFlow is optional")
    def test_tensorflow_graphdef_and_savedmodel(self):
        import tensorflow as tf
        from nn_davinci.adapters.tensorflow import TensorFlowAdapter

        @tf.function(input_signature=[tf.TensorSpec([None, 4], tf.float32)])
        def function(value):
            return tf.nn.relu(value + tf.ones((1, 4)))

        concrete = function.get_concrete_function()
        graph = TensorFlowAdapter().load(concrete).validate()
        self.assertGreater(len(graph.nodes), 2)
        self.assertTrue(graph.edges)
        with tempfile.TemporaryDirectory() as directory:
            module = tf.Module()
            module.serve = function
            tf.saved_model.save(module, directory, signatures={"serving_default": concrete})
            restored = TensorFlowAdapter().load(directory).validate()
            self.assertTrue(restored.nodes)
            self.assertEqual(restored.metadata["source_format"], "tensorflow_graphdef")

    @unittest.skipUnless(importlib.util.find_spec("yaml"), "PyYAML is optional")
    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx is optional")
    def test_yaml_collaboration_and_editable_powerpoint(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            project = Project(
                "team-paper", self.graph,
                collaboration={"members": ["Ada", "Lin"]},
                presentation={"teaching_animation": True, "perspective_mode": True},
            )
            project.add_comment("Check the residual route", author="Ada", target_ids=[self.graph.nodes[2].id])
            restored = Project.load(project.save(root / "paper.yaml"))
            self.assertEqual(restored.comments[0]["author"], "Ada")
            self.assertTrue(restored.presentation["perspective_mode"])
            output = export_graph(
                self.graph, LayoutEngine().layout(self.graph), get_theme(),
                root / "paper", formats=["pptx"],
            )[0]
            deck = Presentation(output)
            self.assertEqual(len(deck.slides), 1)
            self.assertGreater(len(deck.slides[0].shapes), len(self.graph.nodes))

    def test_interactive_html_uses_safe_dom_inspector(self):
        self.graph.nodes[0].name = '<img src=x onerror="alert(1)">'
        html = HtmlRenderer().render(self.graph, LayoutEngine().layout(self.graph), get_theme())
        self.assertIn("info.replaceChildren", html)
        self.assertNotIn("info.innerHTML", html)
        self.assertNotIn('<img src=x onerror="alert(1)">', html)


if __name__ == "__main__":
    unittest.main()
